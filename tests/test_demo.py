import ast
import copy
from datetime import date
import gzip
import hashlib
import json
import os
from pathlib import Path
import re
import unittest
from unittest.mock import patch
import server
from demo_clock import shift_snapshot

ROOT = Path(__file__).resolve().parents[1]


class DemoTests(unittest.TestCase):
    def setUp(self):
        self.env = patch.dict(os.environ, {'RENA_DEMO_DATE': '2026-09-08'})
        self.env.start()
        server.load_day()
        self.client = server.app.test_client()

    def tearDown(self):
        self.env.stop()
        server.load_day()

    def get(self, url):
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200, (url, response.get_data(as_text=True)[:200]))
        data = response.get_json()
        self.assertNotIn('error', data, url)
        return data

    def test_all_main_data_and_calculation_routes(self):
        with patch('socket.create_connection', side_effect=AssertionError('No external connections allowed')):
            for url in ['/api/not-available-fleet/data','/api/reservations','/api/reservations/stats',
                        '/api/fleet-planning?days=7','/api/capacity-control?days=7','/api/island-blocking',
                        '/api/reservations/duplicates','/api/rotation/cargas','/api/rotation/stats',
                        '/api/rotation/historico','/api/infleet/data','/api/match-vehicle?plate=ZZ-01-DA',
                        '/api/chatbot/briefing','/api/demo/info']:
                self.get(url)

    def test_october_opening_preserves_relative_dates_and_metrics(self):
        before = copy.deepcopy(server.calc.SNAPSHOT)
        metrics = self.get('/api/chatbot/briefing')['analysis']
        os.environ['RENA_DEMO_DATE'] = '2026-10-10'
        info = self.get('/api/demo/info')
        self.assertEqual(info['base_date'], '2026-09-08')
        self.assertEqual(info['display_date'], '2026-10-10')
        after = server.calc.SNAPSHOT
        self.assertEqual(len(before['fleet']), len(after['fleet']))
        for old, new in zip(before['reservations'], after['reservations']):
            self.assertEqual(old['RESERVATION_NO'], new['RESERVATION_NO'])
            self.assertEqual((date.fromisoformat(new['PICK_DATETIME'][:10]) - date.fromisoformat(old['PICK_DATETIME'][:10])).days, 32)
            self.assertEqual(old['PICK_DATETIME'][10:], new['PICK_DATETIME'][10:])
        self.assertEqual(self.get('/api/chatbot/briefing')['analysis'], metrics)
        self.assertEqual(self.get('/api/fleet-planning?days=7')['date_from'], '2026-10-10')
        self.assertEqual(self.get('/api/capacity-control?days=7')['days'][0], '2026-10-10')

    def test_rollover_leap_year_and_time_only_slots(self):
        source = {'2028-02-28': ['2028-02-28T23:30:00', '10:00', '08:00–09:00', 'ZZ-01-DA']}
        self.assertEqual(shift_snapshot(source, 1), {'2028-02-29': ['2028-02-29T23:30:00', '10:00', '08:00–09:00', 'ZZ-01-DA']})
        self.assertEqual(shift_snapshot('2026-12-31', 1), '2027-01-01')
        self.assertEqual(shift_snapshot('2028-03-01', -1), '2028-02-29')

    def test_no_source_files_are_changed(self):
        path = ROOT/'data/examples.json.gz'
        original = hashlib.sha256(path.read_bytes()).hexdigest()
        self.get('/api/fleet-planning?days=7')
        self.client.post('/api/chatbot/message', json={'message': 'Saldos Luxury por pool'})
        os.environ['RENA_DEMO_DATE'] = '2027-01-01'
        self.get('/api/chatbot/briefing')
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), original)

    def test_writes_and_extractions_are_rejected(self):
        for path in ['/api/reservations/refresh','/api/not-available-fleet/request-refresh',
                     '/api/infleet/refresh','/api/rotation/cargas','/api/open-email',
                     '/api/permissions','/api/capacity-control/config','/api/not-available-fleet/stations']:
            response = self.client.post(path, json={})
            self.assertEqual(response.status_code, 409, path)
            self.assertFalse(response.get_json()['success'])

    def test_examples_are_synthetic_and_contain_no_real_contact_data(self):
        self.assertTrue(server.ORIGINAL['metadata']['synthetic'])
        self.assertEqual(len(server.ORIGINAL['fleet']), 504)
        self.assertTrue(all(r['_EMAIL'].endswith('@example.invalid') for r in server.ORIGINAL['reservations']))
        self.assertTrue(all(v['licensePlate'].startswith(('ZZ-', 'DEMO-')) for v in server.ORIGINAL['fleet']))
        self.assertTrue(all(c['referencia'].startswith('DEMO-') for c in server.ORIGINAL['rotation']))
        text = json.dumps(server.ORIGINAL)
        self.assertNotRegex(text, r'https?://|[A-Za-z]:\\|\\\\')
        self.assertFalse(list(ROOT.rglob('*.xlsx')))

    def test_html_assets_and_external_connections_policy(self):
        response = self.client.get('/')
        policy = response.headers['Content-Security-Policy']
        self.assertIn("connect-src 'self'", policy)
        for name in ['/vendor/tailwind.js','/vendor/fonts.css','/vendor/leaflet/leaflet.js','/vendor/geography.js',
                     '/chatbot/assets/celebration.svg','/stations_map.html','/chatbot/chatbot.html']:
            r = self.client.get(name)
            self.assertEqual(r.status_code, 200, name)
            r.close()
        response.close()
        for path in (ROOT/'static').rglob('*'):
            if path.suffix in ('.html','.js','.css') and 'vendor' not in path.parts:
                text = path.read_text(encoding='utf-8')
                self.assertNotRegex(text, r'https?://', str(path.relative_to(ROOT)))

    def test_conversation_and_examples_still_have_useful_results(self):
        response = self.client.post('/api/chatbot/message', json={'message': 'Que grupos Luxury ficam negativos por pool?'})
        data = response.get_json()
        self.assertTrue(data['success'])
        self.assertGreater(data['table']['total'], 0)
        follow = self.client.post('/api/chatbot/message', json={'message': 'E amanhã?', 'context': data['context']}).get_json()
        self.assertTrue(follow['success'])
        metrics = self.get('/api/chatbot/briefing')['analysis']
        self.assertGreater(metrics['capacity_exceeded'], 0)
        self.assertGreater(metrics['negative_group_days'], 0)
        self.assertGreater(self.get('/api/reservations/duplicates')['total_groups'], 0)

    def test_internal_files_are_not_served(self):
        for url in ['/server.py','/data/examples.json.gz','/../server.py']:
            self.assertEqual(self.client.get(url).status_code, 404)

    def test_exports_use_fictional_records_and_shifted_dates(self):
        import io
        from openpyxl import load_workbook
        os.environ['RENA_DEMO_DATE'] = '2026-10-10'
        response = self.client.get('/api/reservations/export?format=csv&date_from=2026-10-10&date_to=2026-10-10')
        self.assertEqual(response.status_code, 200)
        self.assertIn('10:00 / 10/10/2026', response.get_data(as_text=True))
        response = self.client.get('/api/reservations/duplicates/csv')
        self.assertEqual(response.status_code, 200)
        self.assertIn('Demo', response.get_data(as_text=True))
        for url in ['/api/reservations/duplicates/excel','/api/reservations/export?format=xlsx']:
            response = self.client.get(url)
            self.assertEqual(response.status_code, 200, url)
            book = load_workbook(io.BytesIO(response.data))
            self.assertGreater(book.active.max_row, 1)
            book.close()


if __name__ == '__main__':
    unittest.main()
