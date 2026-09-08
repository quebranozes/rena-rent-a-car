"""Local calculation environment. No connectors, external paths or background jobs."""
from pathlib import Path
from datetime import timedelta
from demo_clock import DemoDate as date, DemoDateTime as datetime, LISBON
import json
import logging
import os
import re
import time
from flask import Flask, jsonify, request

ROOT = Path(__file__).resolve().parent
app = Flask(__name__, static_folder=None)
app.config['MAX_CONTENT_LENGTH'] = 16384
logger = logging.getLogger('rena-demo')
DATA_DIR = str(ROOT/'data')
STATIONS_FILE = str(ROOT/'data/stations.json')
PLANNING_CONFIG_FILE = str(ROOT/'data/planning_config.json')
CACHE_FILE = str(ROOT/'data/unused.json')
_TZ_LISBON = LISBON
_ACRISS_ALIAS = {}
_plate_overrides = {}
_zone_cache = {'id_zona': {}, 'id_region': {}, 'zona_to_region': {}, 'zonas': [], 'ts': 0}
_RESERVATIONS_SORT_WHITELIST = {'RESERVATION_NO','STATUS','ZGROUP','PICK_DATETIME','RET_DATETIME','DURATION','PICK_STATION_NAME','RET_STATION_NAME','LICENCE_PLATE'}
SNAPSHOT = {}


def configure(snapshot):
    global SNAPSHOT, _not_avail_fleet_cache, _res_bulk_cache, _infleet_cache, _DEMO_ALL_BRANCHES
    SNAPSHOT = snapshot
    _not_avail_fleet_cache = {'data': snapshot['fleet'], 'captured_at': snapshot['captured_at']}
    _res_bulk_cache = {'rows': snapshot['reservations'], 'fetched_at': snapshot['fetched_at']}
    _infleet_cache = {'data': snapshot['infleet']}
    _DEMO_ALL_BRANCHES = snapshot['stations']['stations']
    _zone_cache['ts'] = 0


def _stations_config():
    return SNAPSHOT['stations']


def _read_fleet_type_map(*, read_only=False):
    return SNAPSHOT['types']


def _res_ensure_bulk():
    return None


def _load_capacity_config():
    return SNAPSHOT['capacity']


def _infleet_fetch_status():
    return {'active': False}


def _derive_booking_source(email):
    return 'Demo'


def _rotation_load_cargas():
    return SNAPSHOT['rotation']


def _rotation_load_historico():
    return SNAPSHOT['history']


def _res_station_zone_maps():
    id_zona = {}
    id_region = {}
    zona_to_region = {}
    all_zonas = set()
    try:
        if os.path.exists(STATIONS_FILE):
            cfg = _stations_config()
            pool_region = {}
            for r in cfg.get('planning_regions', []):
                for p in r.get('pools', []):
                    pool_region[p] = r['nome']
            for r in cfg.get('pool_folders', []):
                for p in r.get('pools', []):
                    pool_region[p] = r['nome']
            for sid, sdata in cfg.get('stations', {}).items():
                zona = sdata.get('zona', '') or ''
                region = pool_region.get(zona, '') or ''
                id_zona[sid] = zona
                id_region[sid] = region
                if zona:
                    all_zonas.add(zona)
                    zona_to_region[zona] = region
    except Exception:
        pass
    return (id_zona, id_region, zona_to_region, sorted(all_zonas))

def _get_zone_maps():
    if time.time() - _zone_cache['ts'] > 300:
        iz, ir, zr, zl = _res_station_zone_maps()
        _zone_cache.update(id_zona=iz, id_region=ir, zona_to_region=zr, zonas=zl, ts=time.time())
    return _zone_cache

def _res_filter_rows(rows, args):
    result = list(rows)
    status_filter = args.get('status', '').strip()
    if status_filter:
        statuses = set((s.strip() for s in status_filter.split(',') if s.strip()))
        result = [r for r in result if r.get('STATUS') in statuses]
    pick_station = args.get('pick_station', '').strip()
    if pick_station:
        ps = pick_station.lower()
        result = [r for r in result if r.get('PICK_STATION_NAME') and ps in r['PICK_STATION_NAME'].lower()]
    ret_station = args.get('ret_station', '').strip()
    if ret_station:
        rs = ret_station.lower()
        result = [r for r in result if r.get('RET_STATION_NAME') and rs in r['RET_STATION_NAME'].lower()]
    date_from = args.get('date_from', '').strip()
    date_to = args.get('date_to', '').strip()
    if date_from:
        result = [r for r in result if r.get('PICK_DATETIME') and str(r['PICK_DATETIME']) >= date_from]
    if date_to:
        dt_max = date_to + 'T23:59:59' if 'T' not in date_to else date_to
        result = [r for r in result if r.get('PICK_DATETIME') and str(r['PICK_DATETIME']) <= dt_max]
    search = args.get('search', '').strip()
    if search:
        sl = search.lower()

        def _match(r):
            for f in ('RESERVATION_NO', '_FIRST_NAME', '_LAST_NAME', '_EMAIL'):
                v = r.get(f)
                if v and sl in str(v).lower():
                    return True
            return False
        result = [r for r in result if _match(r)]
    zgroup = args.get('zgroup', '').strip()
    if zgroup:
        wanted_groups = {g.strip() for g in zgroup.split(',') if g.strip()}
        result = [r for r in result if r.get('ZGROUP') in wanted_groups]
    ow_from = args.get('ow_from', '').strip()
    ow_to = args.get('ow_to', '').strip()
    if ow_from or ow_to:
        zc = _get_zone_maps()
        id_zona = zc['id_zona']

        def _ow_match(r):
            if not r.get('ONE_WAY'):
                return False
            pr = id_zona.get(str(r.get('PICK_STATION_ID', '')), '')
            rr = id_zona.get(str(r.get('RET_STATION_ID', '')), '')
            if ow_from and pr != ow_from:
                return False
            if ow_to and rr != ow_to:
                return False
            if pr == rr:
                return False
            return True
        result = [r for r in result if _ow_match(r)]
    foreign = args.get('foreign', '').strip()
    if foreign == '1':
        known_ids = set(_DEMO_ALL_BRANCHES.keys())
        result = [r for r in result if str(r.get('RET_STATION_ID') or '') and str(r.get('RET_STATION_ID')) not in known_ids]
    return result

@app.route('/api/reservations')
def api_reservations():
    try:
        _res_ensure_bulk()
        all_rows = _res_bulk_cache['rows']
        fetched_at = _res_bulk_cache['fetched_at']
        filtered = _res_filter_rows(all_rows, request.args.to_dict())
        unique_nos = set((r.get('RESERVATION_NO') for r in filtered))
        total = len(unique_nos)
        sort_by = request.args.get('sort_by', 'PICK_DATETIME')
        if sort_by not in _RESERVATIONS_SORT_WHITELIST:
            sort_by = 'PICK_DATETIME'
        sort_desc = request.args.get('sort_dir', 'ASC').upper() == 'DESC'
        filtered.sort(key=lambda r: (r.get(sort_by) is None, r.get(sort_by, '')), reverse=sort_desc)
        page = max(1, int(request.args.get('page', 1)))
        per_page = min(200, max(10, int(request.args.get('per_page', 50))))
        offset = (page - 1) * per_page
        page_rows = filtered[offset:offset + per_page]
        zc = _get_zone_maps()
        id_zona = zc['id_zona']
        enriched = []
        id_region = zc['id_region']
        for r in page_rows:
            row = dict(r)
            row['PICK_ZONE'] = id_zona.get(str(r.get('PICK_STATION_ID', '')), '')
            row['RET_ZONE'] = id_zona.get(str(r.get('RET_STATION_ID', '')), '')
            row['PICK_REGION'] = id_region.get(str(r.get('PICK_STATION_ID', '')), '')
            enriched.append(row)
        return jsonify({'rows': enriched, 'total': total, 'page': page, 'per_page': per_page, 'pages': (total + per_page - 1) // per_page if total else 0, 'fetched_at': fetched_at, 'cached': True, 'total_in_cache': len(all_rows)})
    except RuntimeError as e:
        logger.error(f'Reservations query error: {e}')
        return (jsonify({'error': str(e), 'rows': [], 'total': 0, 'page': 1, 'per_page': 50, 'pages': 0, 'fetched_at': None, 'no_token': True}), 200)
    except Exception as e:
        logger.error(f'Reservations query error: {e}')
        return (jsonify({'error': str(e)}), 500)
_RESERVATIONS_EXPORT_COLS = [('RESERVATION_NO', 'Reservation No.'), ('STATUS', 'Status'), ('ZGROUP', 'Group'), ('ACRISS_CODE', 'ACRISS'), ('DURATION', 'Duration'), ('IS_LONGTERM', 'Long-Term'), ('ONE_WAY', 'One Way'), ('IS_PREPAID', 'Prepaid'), ('IS_VIP', 'VIP'), ('LOYALTY_STATUS', 'Loyalty'), ('PICK_DATETIME', 'Pickup Date'), ('PICK_STATION_ID', 'Pickup Station ID'), ('PICK_STATION_NAME', 'Pickup Station'), ('PICK_TIMEZONE', 'Pickup Timezone'), ('RET_DATETIME', 'Return Date'), ('RET_STATION_ID', 'Return Station ID'), ('RET_STATION_NAME', 'Return Station'), ('VEHICLE_TYPE', 'Vehicle Type'), ('VEHICLE_MAKE', 'Manufacturer'), ('VEHICLE_MODEL', 'Model'), ('LICENCE_PLATE', 'License Plate'), ('IS_DELIVERY', 'Delivery'), ('IS_COLLECTION', 'Collection'), ('GAT_DELIVERY', 'GAT Delivery'), ('IS_SUBSCRIPTION', 'Subscription'), ('IS_VEHICLE_EXCHANGE', 'Vehicle Exchange'), ('CHECKOUT_AGENT', 'Checkout Agent'), ('CAN_CHECKOUT', 'Can Check Out')]

def _reservations_export_rows():
    _res_ensure_bulk()
    all_rows = _res_bulk_cache['rows']
    args = request.args.to_dict()
    args['status'] = 'RS,CO'
    filtered = _res_filter_rows(all_rows, args)
    sort_by = request.args.get('sort_by', 'PICK_DATETIME')
    if sort_by not in _RESERVATIONS_SORT_WHITELIST:
        sort_by = 'PICK_DATETIME'
    sort_desc = request.args.get('sort_dir', 'ASC').upper() == 'DESC'
    filtered.sort(key=lambda r: (r.get(sort_by) is None, r.get(sort_by, '')), reverse=sort_desc)
    return filtered

@app.route('/api/reservations/export')
def api_reservations_export():
    try:
        from rena_dates import date_label
        rows = _reservations_export_rows()

        def export_values(row):
            return [date_label(row.get(field), _TZ_LISBON, missing='') if field in ('PICK_DATETIME', 'RET_DATETIME') else row.get(field, '') for field, _label in _RESERVATIONS_EXPORT_COLS]
        fmt = request.args.get('format', 'csv').lower()
        fname = f'reservas_{date.today().isoformat()}'
        if fmt == 'xlsx':
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            import io as io_mod
            wb = Workbook()
            ws = wb.active
            ws.title = 'Reservations'
            header_font = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
            header_fill = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
            header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
            for ci, (_field, label) in enumerate(_RESERVATIONS_EXPORT_COLS, 1):
                cell = ws.cell(row=1, column=ci, value=label)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align
            for r in rows:
                ws.append(export_values(r))
            last_col = ws.cell(row=1, column=len(_RESERVATIONS_EXPORT_COLS)).column_letter
            ws.auto_filter.ref = f'A1:{last_col}1'
            ws.freeze_panes = 'A2'
            for ci in range(1, len(_RESERVATIONS_EXPORT_COLS) + 1):
                ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = 18
            buf = io_mod.BytesIO()
            wb.save(buf)
            return app.response_class(buf.getvalue(), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': f'attachment; filename="{fname}.xlsx"'})
        import csv
        import io as io_mod
        buf = io_mod.StringIO()
        buf.write(chr(65279))
        writer = csv.writer(buf, delimiter=';')
        writer.writerow([label for _field, label in _RESERVATIONS_EXPORT_COLS])
        for r in rows:
            writer.writerow(export_values(r))
        return app.response_class(buf.getvalue().encode('utf-8'), mimetype='text/csv; charset=utf-8', headers={'Content-Disposition': f'attachment; filename="{fname}.csv"'})
    except Exception as e:
        logger.error(f'Reservations export error: {e}')
        return (jsonify({'error': str(e)}), 500)

@app.route('/api/reservations/stats')
def api_reservations_stats():
    try:
        _res_ensure_bulk()
        all_rows = _res_bulk_cache['rows']
        fetched_at = _res_bulk_cache['fetched_at']
        filtered = list(all_rows)
        date_from = request.args.get('date_from', '').strip()
        date_to = request.args.get('date_to', '').strip()
        if date_from:
            filtered = [r for r in filtered if r.get('PICK_DATETIME') and str(r['PICK_DATETIME']) >= date_from]
        if date_to:
            dt_max = date_to + 'T23:59:59' if 'T' not in date_to else date_to
            filtered = [r for r in filtered if r.get('PICK_DATETIME') and str(r['PICK_DATETIME']) <= dt_max]
        from collections import Counter
        status_res = {}
        for r in filtered:
            st = r.get('STATUS')
            if st:
                status_res.setdefault(st, set()).add(r.get('RESERVATION_NO'))
        statuses = {st: len(nos) for st, nos in status_res.items()}
        pick_stations = sorted(set((r['PICK_STATION_NAME'] for r in all_rows if r.get('PICK_STATION_NAME'))))
        zgroups = sorted(set((r['ZGROUP'] for r in all_rows if r.get('ZGROUP'))))
        zc = _get_zone_maps()
        zones = sorted(zc['zonas'])
        return jsonify({'statuses': statuses, 'pick_stations': pick_stations, 'zgroups': zgroups, 'zones': zones, 'fetched_at': fetched_at, 'cached': True})
    except RuntimeError as e:
        logger.error(f'Reservations stats error: {e}')
        return (jsonify({'error': str(e), 'statuses': {}, 'pick_stations': [], 'zgroups': [], 'fetched_at': None, 'no_token': True}), 200)
    except Exception as e:
        logger.error(f'Reservations stats error: {e}')
        return (jsonify({'error': str(e)}), 500)

@app.route('/api/reservations/duplicates')
def api_reservations_duplicates(_assistant=False):
    try:
        if not _assistant:
            _res_ensure_bulk()
        all_rows = [dict(r) for r in _res_bulk_cache.get('rows') or []] if _assistant else _res_bulk_cache['rows']
        fetched_at = _res_bulk_cache['fetched_at']
        today = date.today()
        max_days = int(request.args.get('max_days', 1))
        excl_emails = set()
        excl_names = set()
        try:
            if os.path.exists(STATIONS_FILE):
                cfg = _stations_config()
                for entry in cfg.get('dup_exclusions', []):
                    em = (entry.get('email') or '').strip().lower()
                    nm = (entry.get('nome') or '').strip().lower()
                    if em:
                        excl_emails.add(em)
                    if nm:
                        excl_names.add(nm)
        except Exception as exc:
            logger.warning(f'Could not load dup_exclusions: {exc}')
        from collections import defaultdict
        client_map: dict[tuple, list[dict]] = defaultdict(list)
        for row in all_rows:
            email = row.get('_EMAIL', '')
            fname = row.get('_FIRST_NAME', '')
            lname = row.get('_LAST_NAME', '')
            if not email and (not fname):
                continue
            if email and email.lower() in excl_emails:
                continue
            full_name = f'{fname} {lname}'.strip().lower()
            if full_name and full_name in excl_names:
                continue
            row_status = row.get('STATUS', '')
            if row_status in ('CNL', 'DEN', 'NS', 'PD', 'PP'):
                continue
            key = (email, fname.lower(), lname.lower())
            pick_str = row.get('PICK_DATETIME', '')
            if pick_str:
                try:
                    pick_d = datetime.fromisoformat(pick_str).date()
                    if pick_d < today:
                        continue
                except (ValueError, TypeError):
                    pass
            client_map[key].append(row)

        def _pick_date(row):
            dt_str = row.get('PICK_DATETIME', '')
            if not dt_str:
                return None
            try:
                return datetime.fromisoformat(dt_str).date()
            except (ValueError, TypeError):
                try:
                    return datetime.strptime(dt_str[:10], '%Y-%m-%d').date()
                except Exception:
                    return None

        def _parse_dt(dt_str):
            if not dt_str:
                return None
            try:
                return datetime.fromisoformat(dt_str)
            except (ValueError, TypeError):
                return None

        def _periods_overlap(a, b):
            pick_a = _parse_dt(a.get('PICK_DATETIME', ''))
            ret_a = _parse_dt(a.get('RET_DATETIME', ''))
            pick_b = _parse_dt(b.get('PICK_DATETIME', ''))
            ret_b = _parse_dt(b.get('RET_DATETIME', ''))
            if not pick_a or not pick_b:
                return False
            if not ret_a or not ret_b:
                return pick_a.date() == pick_b.date()
            return pick_a < ret_b and pick_b < ret_a
        groups = []
        for (email, fname_l, lname_l), reservations in client_map.items():
            if len(reservations) < 2:
                continue
            sorted_res = sorted(reservations, key=lambda r: _parse_dt(r.get('PICK_DATETIME', '')) or datetime.min)
            clusters: list[list[dict]] = []
            for row in sorted_res:
                if not _parse_dt(row.get('PICK_DATETIME', '')):
                    continue
                placed = False
                for cluster in clusters:
                    for existing in cluster:
                        if _periods_overlap(row, existing):
                            cluster.append(row)
                            placed = True
                            break
                    if placed:
                        break
                if not placed:
                    clusters.append([row])
            for cluster in clusters:
                if len(cluster) < 2:
                    continue
                groups.append({'client': {'email': email, 'first_name': cluster[0].get('_FIRST_NAME', ''), 'last_name': cluster[0].get('_LAST_NAME', '')}, 'reservations': cluster})
        for group in groups:
            client_email = group['client'].get('email', '')
            source = _derive_booking_source(client_email)
            for r in group['reservations']:
                r['SOURCE'] = source

        def _group_earliest_pick(g):
            dates = [_parse_dt(r.get('PICK_DATETIME', '')) for r in g['reservations']]
            valid = [d.date() for d in dates if d]
            return min(valid) if valid else date.max
        groups.sort(key=_group_earliest_pick)
        return jsonify({'groups': groups, 'total_groups': len(groups), 'total_reservations': sum((len(g['reservations']) for g in groups)), 'fetched_at': fetched_at, 'max_days': max_days})
    except Exception as e:
        logger.error(f'Duplicates query error: {e}')
        return (jsonify({'error': str(e)}), 500)

def _load_planning_config():
    categories = []
    regions = []
    pool_folders = []
    aliases = dict(_ACRISS_ALIAS)
    try:
        if os.path.exists(STATIONS_FILE):
            cfg = _stations_config()
            categories = cfg.get('acriss_categories') or []
            regions = cfg.get('planning_regions') or []
            pool_folders = cfg.get('pool_folders') or []
            for k, v in (cfg.get('acriss_aliases') or {}).items():
                aliases[k.upper()] = v.upper()
    except Exception:
        pass
    return (categories, regions, aliases, pool_folders)

def _resolve_vehicle_branch(v: dict, known_station_ids: set | None=None) -> str:
    if not isinstance(v, dict):
        return ''
    return str((v.get('branch') or {}).get('id') or v.get('currentBranchId') or '')

def _load_station_pool_map() -> dict:
    mapping = {}
    try:
        if os.path.exists(STATIONS_FILE):
            cfg = _stations_config()
            pools = cfg.get('planning_pools') or []
            if pools:
                for pool in pools:
                    nome = pool.get('nome', '')
                    for sid in pool.get('stations', []):
                        mapping[str(sid)] = nome
            else:
                for sid, sdata in cfg.get('stations', {}).items():
                    mapping[sid] = sdata.get('zona', '') or sdata.get('nome', sid)
    except Exception:
        pass
    return mapping

def _occupancy_calc(total: int, unknown: int, avail: int, oficina: int, pronto_of: int, blocked: int, impro: int=0) -> dict:
    base = total - unknown
    parados = avail + oficina + pronto_of + blocked
    em_andamento = base - parados - impro
    pct = round(em_andamento / base * 100, 1) if base > 0 else None
    impro_pct = round(impro / base * 100, 1) if base > 0 else None
    return {'total': total, 'unknown': unknown, 'base': base, 'parados': parados, 'impro': impro, 'impro_pct': impro_pct, 'em_andamento': em_andamento, 'pct': pct}

@app.route('/api/fleet-planning')
def api_fleet_planning(_assistant=False):
    try:
        pool_filter = request.args.get('pool', '').strip()
        days = min(int(request.args.get('days', 30)), 365)
        station_pool = _load_station_pool_map()
        acriss_categories, planning_regions, acriss_aliases, pool_folders = _load_planning_config()
        station_names = {}
        station_tipos = {}
        station_paired = {}
        try:
            if os.path.exists(STATIONS_FILE):
                _stn_cfg = _stations_config()
                for _sid, _sd in _stn_cfg.get('stations', {}).items():
                    station_names[_sid] = _sd.get('nome', _sid)
                    station_tipos[_sid] = _sd.get('tipo', '')
                    if _sd.get('paired'):
                        station_paired[_sid] = str(_sd['paired'])
        except Exception:
            pass
        fleet_vehicles = _not_avail_fleet_cache.get('data') or []
        fleet_captured = _not_avail_fleet_cache.get('captured_at', '')
        _AVAIL_STATUSES = {1, 2, 3, 4, 7, 9, 10, 11}
        _OFICINA_STATUSES = {5, 6}
        _DEFLEET_MARGIN_DAYS = 0
        _IMPRO_RATE_CODES = {'PTCCC', 'PTCCF', 'PTCCS', 'PTCCR', 'PTCCT'}
        _IMPRO_RATE_LABELS = {'PTCCC': 'Courtesy', 'PTCCF': 'Internal Movement / Transport', 'PTCCS': 'Trailer / Trailer Islands', 'PTCCR': 'Workshop / Mechanical Repair / Collision / Reconditioning', 'PTCCT': 'Support Vehicle'}
        fleet_by_pool_group: dict[str, dict[str, int]] = {}
        fleet_total_by_pool: dict[str, int] = {}
        fleet_total_by_pool_group: dict[str, dict[str, int]] = {}
        fleet_oficina_by_pool: dict[str, int] = {}
        fleet_oficina_by_pool_group: dict[str, dict[str, int]] = {}
        unknown_status_by_pool: dict[str, int] = {}
        unknown_status_by_pool_group: dict[str, dict[str, int]] = {}
        impro_by_pool: dict[str, int] = {}
        impro_by_pool_group: dict[str, dict[str, int]] = {}
        impro_vehicles_by_pool: dict[str, list] = {}
        impro_vehicles_by_stn: dict[str, list] = {}
        all_pools_set: set[str] = set()
        today = date.today()
        end_date = today + timedelta(days=days)
        date_range = [today + timedelta(days=d) for d in range(days + 1)]
        planning: dict[str, dict[str, dict]] = {}

        def _ensure_pool(pool_name):
            if pool_name not in planning:
                all_pools_set.add(pool_name)
                planning[pool_name] = {}
                for d in date_range:
                    ds = d.isoformat()
                    planning[pool_name][ds] = {'res_out': 0, 'res_out_lt': 0, 'res_in': 0, 'res_in_lt': 0, 'fleet_in': 0, 'fleet_in_lt': 0, 'defleet_out': 0, 'rented_defleet': 0, 'repair_in': 0, 'groups_out': {}, 'groups_out_lt': {}, 'groups_in': {}, 'groups_in_lt': {}, 'groups_fleet_in': {}, 'groups_fleet_in_lt': {}, 'groups_defleet_out': {}, 'groups_rented_defleet': {}, 'groups_repair_in': {}, 'res_out_details': [], 'res_in_details': [], 'fleet_in_details': [], 'defleet_details': [], 'repair_in_details': []}
        avail_real_by_pool: dict[str, int] = {}
        avail_real_by_pool_group: dict[str, dict[str, int]] = {}
        _STN = '_s:'
        all_stations_set: set[str] = set()

        def _ensure_station(station_id):
            stn_key = _STN + station_id
            if stn_key not in planning:
                all_stations_set.add(station_id)
                planning[stn_key] = {}
                for d in date_range:
                    ds2 = d.isoformat()
                    planning[stn_key][ds2] = {'res_out': 0, 'res_out_lt': 0, 'res_in': 0, 'res_in_lt': 0, 'fleet_in': 0, 'fleet_in_lt': 0, 'defleet_out': 0, 'rented_defleet': 0, 'repair_in': 0, 'groups_out': {}, 'groups_out_lt': {}, 'groups_in': {}, 'groups_in_lt': {}, 'groups_fleet_in': {}, 'groups_fleet_in_lt': {}, 'groups_defleet_out': {}, 'groups_rented_defleet': {}, 'groups_repair_in': {}, 'res_out_details': [], 'res_in_details': [], 'fleet_in_details': [], 'defleet_details': [], 'repair_in_details': []}
        _SKIP_TIPOS = {'Roubo', 'Testes'}
        for _sid, _tipo in station_tipos.items():
            if _tipo not in _SKIP_TIPOS and _sid in station_pool:
                _ensure_pool(station_pool[_sid])
                _ensure_station(_sid)
        avail_real_by_stn: dict[str, int] = {}
        avail_real_by_stn_group: dict[str, dict[str, int]] = {}
        fleet_total_by_stn: dict[str, int] = {}
        occupancy_counts_by_stn: dict = {}
        fleet_oficina_by_stn: dict[str, int] = {}
        fleet_oficina_by_stn_group: dict[str, dict[str, int]] = {}
        fleet_pronto_of_by_pool: dict[str, int] = {}
        fleet_pronto_of_by_pool_group: dict[str, dict[str, int]] = {}
        fleet_pronto_of_by_stn: dict[str, int] = {}
        fleet_pronto_of_by_stn_group: dict[str, dict[str, int]] = {}
        avail_vehicles_by_pool: dict[str, list] = {}
        avail_vehicles_by_stn: dict[str, list] = {}
        oficina_vehicles_by_pool: dict[str, list] = {}
        oficina_vehicles_by_stn: dict[str, list] = {}
        pronto_of_vehicles_by_pool: dict[str, list] = {}
        pronto_of_vehicles_by_stn: dict[str, list] = {}
        blocked_vehicles_by_pool: dict[str, list] = {}
        blocked_vehicles_by_stn: dict[str, list] = {}
        blocked_count_by_pool: dict[str, int] = {}
        blocked_count_by_stn: dict[str, int] = {}
        blocked_groups_by_pool: dict[str, dict[str, int]] = {}
        blocked_groups_by_stn: dict[str, dict[str, int]] = {}
        noshow_by_pool: dict[str, list] = {}
        noshow_by_stn: dict[str, list] = {}
        noshow_groups_by_pool: dict[str, dict[str, int]] = {}
        overdue_ret_by_pool: dict[str, list] = {}
        overdue_ret_groups_by_pool: dict[str, dict[str, int]] = {}
        now_dt = datetime.now()
        tipo_map = _read_fleet_type_map(read_only=_assistant)
        _plate_best: dict[str, dict] = {}
        _ACTIVE_CS = {0, 1, 2, 3, 4, 5, 6, 7, 9, 10, 11}
        for v in fleet_vehicles:
            plate = v.get('licensePlate', '')
            if not plate:
                continue
            override_id = _plate_overrides.get(plate)
            if override_id is not None:
                cur_id = str(v.get('internalNumber', ''))
                if cur_id != str(override_id):
                    continue
            prev = _plate_best.get(plate)
            if prev is None:
                _plate_best[plate] = v
            else:
                prev_cs = prev.get('classicStatus')
                cur_cs = v.get('classicStatus')
                prev_sub = (prev.get('subStatus') or '').upper()
                cur_sub = (v.get('subStatus') or '').upper()
                prev_defleet = 'DEFLEET' in prev_sub or prev_cs not in _ACTIVE_CS
                cur_defleet = 'DEFLEET' in cur_sub or cur_cs not in _ACTIVE_CS
                if prev_defleet and (not cur_defleet):
                    _plate_best[plate] = v
                elif not prev_defleet and cur_defleet:
                    pass
                else:
                    prev_km = prev.get('mileage') or 0
                    cur_km = v.get('mileage') or 0
                    if cur_km > prev_km:
                        _plate_best[plate] = v
        fleet_vehicles_deduped = list(_plate_best.values())
        _seen_plates: set[str] = set()
        for v in fleet_vehicles_deduped:
            plate = v.get('licensePlate', '')
            if plate and plate in _seen_plates:
                continue
            if plate:
                _seen_plates.add(plate)
            branch = _resolve_vehicle_branch(v, set(station_names.keys()))
            pool = station_pool.get(branch)
            if not pool:
                continue
            if pool_filter and pool != pool_filter:
                continue
            acriss = (v.get('acrissCode') or '')[:4].upper()
            acriss = acriss_aliases.get(acriss, acriss)
            all_pools_set.add(pool)
            fleet_total_by_pool.setdefault(pool, 0)
            fleet_total_by_pool[pool] += 1
            fleet_total_by_pool_group.setdefault(pool, {})
            fleet_total_by_pool_group[pool][acriss] = fleet_total_by_pool_group[pool].get(acriss, 0) + 1
            fleet_total_by_stn.setdefault(branch, 0)
            fleet_total_by_stn[branch] += 1
            cs = v.get('classicStatus')
            stn_occ = occupancy_counts_by_stn.setdefault(branch, {}).setdefault(acriss, {'total': 0, 'unknown': 0, 'impro': 0})
            stn_occ['total'] += 1
            if cs not in _AVAIL_STATUSES and cs not in _OFICINA_STATUSES and (cs != 0):
                stn_occ['unknown'] += 1
            if cs == 0 and str((v.get('rentalActivity') or {}).get('rateCode') or '')[:5] in _IMPRO_RATE_CODES:
                stn_occ['impro'] += 1
            if cs not in _AVAIL_STATUSES and cs not in _OFICINA_STATUSES and (cs != 0):
                unknown_status_by_pool[pool] = unknown_status_by_pool.get(pool, 0) + 1
                unknown_status_by_pool_group.setdefault(pool, {})
                unknown_status_by_pool_group[pool][acriss] = unknown_status_by_pool_group[pool].get(acriss, 0) + 1
                continue
            if cs == 0:
                _ra = v.get('rentalActivity') or {}
                rate_code = _ra.get('rateCode', '') or ''
                rate_prefix = rate_code[:5]
                if rate_prefix in _IMPRO_RATE_CODES:
                    impro_by_pool[pool] = impro_by_pool.get(pool, 0) + 1
                    impro_by_pool_group.setdefault(pool, {})
                    impro_by_pool_group[pool][acriss] = impro_by_pool_group[pool].get(acriss, 0) + 1
                    impro_detail = {'plate': v.get('licensePlate', ''), 'name': v.get('displayName', ''), 'group': acriss, 'rate_code': rate_code, 'rate_label': _IMPRO_RATE_LABELS.get(rate_prefix, rate_code), 'station': station_names.get(branch, branch), 'station_id': branch, 'return_time': v.get('returnTime') or '', 'ra_number': str(_ra.get('rentalAgreementNumber', '')), 'mileage': v.get('mileage'), 'engine': v.get('engineType', '')}
                    impro_vehicles_by_pool.setdefault(pool, []).append(impro_detail)
                    impro_vehicles_by_stn.setdefault(branch, []).append(impro_detail)
            is_pronto_of = cs in _AVAIL_STATUSES and station_tipos.get(branch) == 'Workshop'
            if is_pronto_of:
                cs = 5
            if is_pronto_of or cs in _OFICINA_STATUSES:
                oficina_detail = {'plate': v.get('licensePlate', ''), 'name': v.get('displayName', ''), 'group': acriss, 'status': cs, 'pronto_of': is_pronto_of, 'activity': v.get('classicActivityCode', ''), 'station': station_names.get(branch, branch), 'station_id': branch, 'parking': v.get('parkingSlot') or v.get('parkingSlotV2') or v.get('parkingParameterV2') or v.get('parkingParameter') or '', 'in_status_since': v.get('inStatusSince', ''), 'remaining_km': v.get('remainingMileage'), 'mileage': v.get('mileage'), 'defleet': v.get('defleetDate') or '', 'engine': v.get('engineType', ''), 'tipo': tipo_map.get((v.get('licensePlate') or '').upper(), {}).get('tipo_compra', '')}
                if is_pronto_of:
                    fleet_pronto_of_by_pool.setdefault(pool, 0)
                    fleet_pronto_of_by_pool[pool] += 1
                    fleet_pronto_of_by_pool_group.setdefault(pool, {})
                    fleet_pronto_of_by_pool_group[pool][acriss] = fleet_pronto_of_by_pool_group[pool].get(acriss, 0) + 1
                    fleet_pronto_of_by_stn.setdefault(branch, 0)
                    fleet_pronto_of_by_stn[branch] += 1
                    fleet_pronto_of_by_stn_group.setdefault(branch, {})
                    fleet_pronto_of_by_stn_group[branch][acriss] = fleet_pronto_of_by_stn_group[branch].get(acriss, 0) + 1
                    pronto_of_vehicles_by_pool.setdefault(pool, []).append(oficina_detail)
                    pronto_of_vehicles_by_stn.setdefault(branch, []).append(oficina_detail)
                else:
                    fleet_oficina_by_pool.setdefault(pool, 0)
                    fleet_oficina_by_pool[pool] += 1
                    fleet_oficina_by_pool_group.setdefault(pool, {})
                    fleet_oficina_by_pool_group[pool][acriss] = fleet_oficina_by_pool_group[pool].get(acriss, 0) + 1
                    fleet_oficina_by_stn.setdefault(branch, 0)
                    fleet_oficina_by_stn[branch] += 1
                    fleet_oficina_by_stn_group.setdefault(branch, {})
                    fleet_oficina_by_stn_group[branch][acriss] = fleet_oficina_by_stn_group[branch].get(acriss, 0) + 1
                    oficina_vehicles_by_pool.setdefault(pool, []).append(oficina_detail)
                    oficina_vehicles_by_stn.setdefault(branch, []).append(oficina_detail)
                continue
            defleet_str = v.get('defleetDate') or ''
            defleet_d = None
            if defleet_str:
                try:
                    defleet_d = date.fromisoformat(defleet_str)
                except (ValueError, TypeError):
                    pass
            near_defleet = defleet_d and (defleet_d - today).days <= _DEFLEET_MARGIN_DAYS
            if cs in _AVAIL_STATUSES:
                fleet_by_pool_group.setdefault(pool, {})
                fleet_by_pool_group[pool].setdefault(acriss, 0)
                fleet_by_pool_group[pool][acriss] += 1
                _hold_desc = v.get('holdDescriptionForDueDate') or ''
                _hold_code = v.get('holdCodeForDueDate') or ''
                has_alarm = bool(_hold_desc or _hold_code)
                hold_remaining = v.get('holdRemainingMileage')
                if hold_remaining is None:
                    hold_remaining = v.get('remainingMileage')
                low_remaining_km = not has_alarm and hold_remaining is not None and (hold_remaining < 150)
                past_defleet = defleet_d is not None and defleet_d <= today
                if has_alarm or past_defleet or low_remaining_km:
                    reason = 'bloqueado' if has_alarm else 'defleet' if past_defleet else 'sem km'
                    blocked_detail = {'plate': v.get('licensePlate', ''), 'name': v.get('displayName', ''), 'group': acriss, 'status': cs, 'reason': reason, 'remaining_km': hold_remaining, 'mileage': v.get('mileage'), 'defleet': defleet_str, 'station': station_names.get(branch, branch), 'station_id': branch, 'parking': v.get('parkingSlot') or v.get('parkingSlotV2') or v.get('parkingParameterV2') or v.get('parkingParameter') or '', 'engine': v.get('engineType', ''), 'hold_code': v.get('holdCodeForDueDate', ''), 'hold_desc': v.get('holdDescriptionForDueDate', ''), 'sub_status': v.get('subStatus', ''), 'tipo': tipo_map.get((v.get('licensePlate') or '').upper(), {}).get('tipo_compra', '')}
                    blocked_vehicles_by_pool.setdefault(pool, []).append(blocked_detail)
                    blocked_vehicles_by_stn.setdefault(branch, []).append(blocked_detail)
                    blocked_count_by_pool[pool] = blocked_count_by_pool.get(pool, 0) + 1
                    blocked_count_by_stn[branch] = blocked_count_by_stn.get(branch, 0) + 1
                    blocked_groups_by_pool.setdefault(pool, {})
                    blocked_groups_by_pool[pool][acriss] = blocked_groups_by_pool[pool].get(acriss, 0) + 1
                    blocked_groups_by_stn.setdefault(branch, {})
                    blocked_groups_by_stn[branch][acriss] = blocked_groups_by_stn[branch].get(acriss, 0) + 1
                else:
                    avail_real_by_pool.setdefault(pool, 0)
                    avail_real_by_pool[pool] += 1
                    avail_real_by_pool_group.setdefault(pool, {})
                    avail_real_by_pool_group[pool].setdefault(acriss, 0)
                    avail_real_by_pool_group[pool][acriss] += 1
                    avail_real_by_stn.setdefault(branch, 0)
                    avail_real_by_stn[branch] += 1
                    avail_real_by_stn_group.setdefault(branch, {})
                    avail_real_by_stn_group[branch].setdefault(acriss, 0)
                    avail_real_by_stn_group[branch][acriss] += 1
                    veh_detail = {'plate': v.get('licensePlate', ''), 'name': v.get('displayName', ''), 'group': acriss, 'status': cs, 'parking': v.get('parkingSlot') or v.get('parkingSlotV2') or v.get('parkingParameterV2') or v.get('parkingParameter') or '', 'station': station_names.get(branch, branch), 'station_id': branch, 'remaining_km': hold_remaining, 'mileage': v.get('mileage'), 'defleet': defleet_str, 'fuel': v.get('fuelLevel'), 'charge_pct': v.get('chargeLevelPercentage'), 'engine': v.get('engineType', ''), 'tipo': tipo_map.get((v.get('licensePlate') or '').upper(), {}).get('tipo_compra', '')}
                    if near_defleet:
                        veh_detail['near_defleet'] = True
                    avail_vehicles_by_pool.setdefault(pool, []).append(veh_detail)
                    avail_vehicles_by_stn.setdefault(branch, []).append(veh_detail)
                if defleet_d and today <= defleet_d <= end_date:
                    _ensure_pool(pool)
                    _ensure_station(branch)
                    ds = defleet_d.isoformat()
                    defleet_detail = {'plate': v.get('licensePlate', ''), 'name': v.get('displayName', ''), 'group': acriss, 'station': station_names.get(branch, branch), 'station_id': branch, 'defleet_date': defleet_str, 'remaining_km': hold_remaining, 'mileage': v.get('mileage'), 'remaining_days': (defleet_d - today).days, 'engine': v.get('engineType', '')}
                    if ds in planning[pool]:
                        if not past_defleet:
                            planning[pool][ds]['defleet_out'] += 1
                            planning[pool][ds]['groups_defleet_out'][acriss] = planning[pool][ds]['groups_defleet_out'].get(acriss, 0) + 1
                        planning[pool][ds]['defleet_details'].append(defleet_detail)
                    stn_key = _STN + branch
                    if ds in planning[stn_key]:
                        if not past_defleet:
                            planning[stn_key][ds]['defleet_out'] += 1
                            planning[stn_key][ds]['groups_defleet_out'][acriss] = planning[stn_key][ds]['groups_defleet_out'].get(acriss, 0) + 1
                        planning[stn_key][ds]['defleet_details'].append(defleet_detail)
            elif cs == 0:
                ret_time = v.get('returnTime') or ''
                if not ret_time:
                    continue
                try:
                    ret_d = datetime.fromisoformat(ret_time.replace('Z', '+00:00')).date()
                except (ValueError, TypeError):
                    continue
                if ret_d < today or ret_d > end_date:
                    if ret_d < today:
                        _ret_branch_raw_ov = v.get('returnBranch') or {}
                        _ret_branch_id_ov = str(_ret_branch_raw_ov.get('id', '') if isinstance(_ret_branch_raw_ov, dict) else _ret_branch_raw_ov or '')
                        ret_bid_ov = _ret_branch_id_ov if _ret_branch_id_ov else str(v.get('currentBranchId') or v.get('branch', {}).get('id', '') or '')
                        ret_pool_ov = station_pool.get(ret_bid_ov)
                        if ret_pool_ov and (not pool_filter or ret_pool_ov == pool_filter):
                            _ov_is_lt = (v.get('rentalActivity') or {}).get('isLongterm', False)
                            _ov_ra = v.get('rentalActivity') or {}
                            _ov_detail = {'plate': v.get('licensePlate', ''), 'name': v.get('displayName', ''), 'group': acriss, 'station': station_names.get(ret_bid_ov, ret_bid_ov), 'station_id': ret_bid_ov, 'ret_time': ret_time, 'days_overdue': (today - ret_d).days, 'is_lt': _ov_is_lt, 'ra_number': str(_ov_ra.get('rentalAgreementNumber', ''))}
                            overdue_ret_by_pool.setdefault(ret_pool_ov, []).append(_ov_detail)
                            _ov_grp = overdue_ret_groups_by_pool.setdefault(ret_pool_ov, {})
                            _ov_grp[acriss] = _ov_grp.get(acriss, 0) + 1
                    continue
                _ret_branch_raw = v.get('returnBranch') or {}
                _ret_branch_id = str(_ret_branch_raw.get('id', '') if isinstance(_ret_branch_raw, dict) else _ret_branch_raw or '')
                ret_bid = _ret_branch_id if _ret_branch_id else str(v.get('currentBranchId') or v.get('branch', {}).get('id', '') or '')
                ret_pool = station_pool.get(ret_bid)
                if not ret_pool:
                    continue
                if pool_filter and ret_pool != pool_filter:
                    continue
                will_defleet = defleet_d and (defleet_d - ret_d).days <= _DEFLEET_MARGIN_DAYS
                is_lt = (v.get('rentalActivity') or {}).get('isLongterm', False)
                _KM_PER_DAY = 150
                remaining_km_now = v.get('holdRemainingMileage')
                if remaining_km_now is None:
                    remaining_km_now = v.get('remainingMileage')
                days_to_return = max((ret_d - today).days, 1)
                predicted_km_at_return = None
                will_exceed_km = False
                if remaining_km_now is not None:
                    predicted_km_at_return = remaining_km_now - days_to_return * _KM_PER_DAY
                    if predicted_km_at_return < 0:
                        will_exceed_km = True
                _ensure_pool(ret_pool)
                _ensure_station(ret_bid)
                ds = ret_d.isoformat()
                stn_key = _STN + ret_bid
                ra = v.get('rentalActivity') or {}
                _fleet_plate = v.get('licensePlate', '')
                _pt_plate_re = '^[A-Z0-9]{2}-[A-Z0-9]{2}-[A-Z0-9]{2}$'
                _is_foreign_plate = not bool(re.match(_pt_plate_re, _fleet_plate.strip().upper())) if _fleet_plate.strip() else False
                fleet_ret_detail = {'plate': _fleet_plate, 'name': v.get('displayName', ''), 'group': acriss, 'station': station_names.get(ret_bid, ret_bid), 'ret_time': ret_time, 'is_lt': is_lt, 'will_defleet': bool(will_defleet), 'will_exceed_km': will_exceed_km, 'remaining_km': remaining_km_now, 'predicted_km': predicted_km_at_return, 'days_to_return': days_to_return, 'ra_number': str(ra.get('rentalAgreementNumber', '')), 'is_foreign': _is_foreign_plate}
                if will_defleet or (will_exceed_km and (not is_lt)):
                    rented_defleet_detail = {'plate': v.get('licensePlate', ''), 'name': v.get('displayName', ''), 'group': acriss, 'station': station_names.get(ret_bid, ret_bid), 'station_id': ret_bid, 'defleet_date': defleet_str, 'remaining_km': v.get('holdRemainingMileage') or v.get('remainingMileage'), 'predicted_km': predicted_km_at_return, 'mileage': v.get('mileage'), 'remaining_days': (defleet_d - today).days if defleet_d else None, 'engine': v.get('engineType', ''), 'reason': 'defleet' if will_defleet else 'sem_km', 'ret_time': ret_time}
                is_ret_repair = station_tipos.get(ret_bid) == 'Workshop'
                if ds in planning[ret_pool]:
                    if will_defleet:
                        planning[ret_pool][ds]['rented_defleet'] += 1
                        planning[ret_pool][ds]['groups_rented_defleet'][acriss] = planning[ret_pool][ds]['groups_rented_defleet'].get(acriss, 0) + 1
                        planning[ret_pool][ds]['defleet_details'].append(rented_defleet_detail)
                    elif will_exceed_km and (not is_lt):
                        planning[ret_pool][ds]['rented_defleet'] += 1
                        planning[ret_pool][ds]['groups_rented_defleet'][acriss] = planning[ret_pool][ds]['groups_rented_defleet'].get(acriss, 0) + 1
                        planning[ret_pool][ds]['defleet_details'].append(rented_defleet_detail)
                    elif is_ret_repair:
                        planning[ret_pool][ds]['repair_in'] += 1
                        planning[ret_pool][ds]['groups_repair_in'][acriss] = planning[ret_pool][ds]['groups_repair_in'].get(acriss, 0) + 1
                        planning[ret_pool][ds]['repair_in_details'].append(dict(fleet_ret_detail, type='fleet'))
                    elif _is_foreign_plate:
                        pass
                    elif is_lt:
                        planning[ret_pool][ds]['fleet_in_lt'] += 1
                        planning[ret_pool][ds]['groups_fleet_in_lt'][acriss] = planning[ret_pool][ds]['groups_fleet_in_lt'].get(acriss, 0) + 1
                    else:
                        planning[ret_pool][ds]['fleet_in'] += 1
                        planning[ret_pool][ds]['groups_fleet_in'][acriss] = planning[ret_pool][ds]['groups_fleet_in'].get(acriss, 0) + 1
                    if not is_ret_repair:
                        planning[ret_pool][ds]['fleet_in_details'].append(fleet_ret_detail)
                if ds in planning[stn_key]:
                    if will_defleet:
                        planning[stn_key][ds]['rented_defleet'] += 1
                        planning[stn_key][ds]['groups_rented_defleet'][acriss] = planning[stn_key][ds]['groups_rented_defleet'].get(acriss, 0) + 1
                        planning[stn_key][ds]['defleet_details'].append(rented_defleet_detail)
                    elif will_exceed_km and (not is_lt):
                        planning[stn_key][ds]['rented_defleet'] += 1
                        planning[stn_key][ds]['groups_rented_defleet'][acriss] = planning[stn_key][ds]['groups_rented_defleet'].get(acriss, 0) + 1
                        planning[stn_key][ds]['defleet_details'].append(rented_defleet_detail)
                    elif is_ret_repair:
                        planning[stn_key][ds]['repair_in'] += 1
                        planning[stn_key][ds]['groups_repair_in'][acriss] = planning[stn_key][ds]['groups_repair_in'].get(acriss, 0) + 1
                        planning[stn_key][ds]['repair_in_details'].append(dict(fleet_ret_detail, type='fleet'))
                    elif _is_foreign_plate:
                        pass
                    elif is_lt:
                        planning[stn_key][ds]['fleet_in_lt'] += 1
                        planning[stn_key][ds]['groups_fleet_in_lt'][acriss] = planning[stn_key][ds]['groups_fleet_in_lt'].get(acriss, 0) + 1
                    else:
                        planning[stn_key][ds]['fleet_in'] += 1
                        planning[stn_key][ds]['groups_fleet_in'][acriss] = planning[stn_key][ds]['groups_fleet_in'].get(acriss, 0) + 1
                    if not is_ret_repair:
                        planning[stn_key][ds]['fleet_in_details'].append(fleet_ret_detail)
                if not will_defleet and (not is_lt) and (not is_ret_repair) and defleet_d:
                    if today <= defleet_d <= end_date:
                        _ensure_pool(ret_pool)
                        ds2 = defleet_d.isoformat()
                        future_defleet_detail = {'plate': v.get('licensePlate', ''), 'name': v.get('displayName', ''), 'group': acriss, 'station': station_names.get(ret_bid, ret_bid), 'station_id': ret_bid, 'defleet_date': defleet_str, 'remaining_km': v.get('holdRemainingMileage') or v.get('remainingMileage'), 'mileage': v.get('mileage'), 'remaining_days': (defleet_d - today).days, 'engine': v.get('engineType', '')}
                        if ds2 in planning[ret_pool]:
                            planning[ret_pool][ds2]['defleet_out'] += 1
                            planning[ret_pool][ds2]['groups_defleet_out'][acriss] = planning[ret_pool][ds2]['groups_defleet_out'].get(acriss, 0) + 1
                            planning[ret_pool][ds2]['defleet_details'].append(future_defleet_detail)
                        if ds2 in planning[stn_key]:
                            planning[stn_key][ds2]['defleet_out'] += 1
                            planning[stn_key][ds2]['groups_defleet_out'][acriss] = planning[stn_key][ds2]['groups_defleet_out'].get(acriss, 0) + 1
                            planning[stn_key][ds2]['defleet_details'].append(future_defleet_detail)
        all_res = _res_bulk_cache.get('rows') or []
        res_fetched_at = _res_bulk_cache.get('fetched_at', '')
        res_total_cache = len(all_res)
        res_counted = 0
        for row in all_res:
            zgroup = row.get('ZGROUP', '')
            zgroup = acriss_aliases.get(zgroup, zgroup)
            is_lt = row.get('IS_LONGTERM', False)
            status = row.get('STATUS', '')
            if status not in ('RS', 'RQ', 'ARQ'):
                continue
            pick_dt = row.get('PICK_DATETIME', '')
            _is_noshow = False
            if pick_dt:
                try:
                    pick_datetime = datetime.fromisoformat(pick_dt)
                    if pick_datetime.tzinfo is not None:
                        pick_datetime = pick_datetime.replace(tzinfo=None)
                    pick_date = pick_datetime.date()
                except (ValueError, TypeError):
                    pick_datetime = None
                    pick_date = None
                if pick_date and pick_datetime:
                    if pick_date < today:
                        _is_noshow = True
                    elif pick_date == today:
                        if pick_datetime + timedelta(hours=1) < now_dt:
                            _is_noshow = True
            pick_sid = str(row.get('PICK_STATION_ID', ''))
            pick_pool = station_pool.get(pick_sid)
            if not pick_pool:
                pick_pool = None
            elif pool_filter and pick_pool != pool_filter:
                pick_pool = None
            res_detail = {'res_no': row.get('RESERVATION_NO', ''), 'group': zgroup, 'pick_station': station_names.get(pick_sid, row.get('PICK_STATION_NAME', pick_sid)), 'ret_station': station_names.get(str(row.get('RET_STATION_ID', '')), row.get('RET_STATION_NAME', '')), 'pick_time': pick_dt[:16].replace('T', ' ') if pick_dt and len(pick_dt) >= 16 else pick_dt, 'ret_time': (row.get('RET_DATETIME', '') or '')[:16].replace('T', ' '), 'duration': row.get('DURATION', ''), 'is_lt': is_lt, 'one_way': row.get('ONE_WAY', False), 'status': status, 'source': '', 'is_vip': row.get('IS_VIP', False), 'is_prepaid': row.get('IS_PREPAID', False), 'loyalty': row.get('LOYALTY_STATUS', '') or '', 'plate': (row.get('LICENCE_PLATE') or '').strip()}
            if _is_noshow:
                if pick_pool and pick_date == today:
                    noshow_by_pool.setdefault(pick_pool, []).append(res_detail)
                    _nsg = noshow_groups_by_pool.setdefault(pick_pool, {})
                    _nsg[zgroup] = _nsg.get(zgroup, 0) + 1
                    noshow_by_stn.setdefault(pick_sid, []).append(res_detail)
                continue
            if pick_dt and pick_pool:
                try:
                    pd = datetime.fromisoformat(pick_dt).date()
                except (ValueError, TypeError):
                    pd = None
                if pd and today <= pd <= end_date:
                    _ensure_pool(pick_pool)
                    _ensure_station(pick_sid)
                    ds = pd.isoformat()
                    if ds in planning[pick_pool]:
                        planning[pick_pool][ds]['res_out'] += 1
                        planning[pick_pool][ds]['groups_out'][zgroup] = planning[pick_pool][ds]['groups_out'].get(zgroup, 0) + 1
                        planning[pick_pool][ds]['res_out_details'].append(res_detail)
                        res_counted += 1
                    stn_key = _STN + pick_sid
                    if ds in planning[stn_key]:
                        planning[stn_key][ds]['res_out'] += 1
                        planning[stn_key][ds]['groups_out'][zgroup] = planning[stn_key][ds]['groups_out'].get(zgroup, 0) + 1
                        planning[stn_key][ds]['res_out_details'].append(res_detail)
            ret_dt = row.get('RET_DATETIME', '')
            ret_sid = str(row.get('RET_STATION_ID', ''))
            ret_pool = station_pool.get(ret_sid)
            if not ret_pool:
                ret_pool = None
            elif pool_filter and ret_pool != pool_filter:
                ret_pool = None
            if ret_dt and ret_pool:
                try:
                    rd = datetime.fromisoformat(ret_dt).date()
                except (ValueError, TypeError):
                    rd = None
                if rd and today <= rd <= end_date:
                    _ensure_pool(ret_pool)
                    _ensure_station(ret_sid)
                    ds = rd.isoformat()
                    is_ret_repair = station_tipos.get(ret_sid) == 'Workshop'
                    if ds in planning[ret_pool]:
                        if is_ret_repair:
                            planning[ret_pool][ds]['repair_in'] += 1
                            planning[ret_pool][ds]['groups_repair_in'][zgroup] = planning[ret_pool][ds]['groups_repair_in'].get(zgroup, 0) + 1
                            planning[ret_pool][ds]['repair_in_details'].append(dict(res_detail, type='res'))
                        elif is_lt:
                            planning[ret_pool][ds]['res_in_lt'] += 1
                            planning[ret_pool][ds]['groups_in_lt'][zgroup] = planning[ret_pool][ds]['groups_in_lt'].get(zgroup, 0) + 1
                            planning[ret_pool][ds]['res_in_details'].append(res_detail)
                        else:
                            planning[ret_pool][ds]['res_in'] += 1
                            planning[ret_pool][ds]['groups_in'][zgroup] = planning[ret_pool][ds]['groups_in'].get(zgroup, 0) + 1
                            planning[ret_pool][ds]['res_in_details'].append(res_detail)
                    stn_key = _STN + ret_sid
                    if ds in planning[stn_key]:
                        if is_ret_repair:
                            planning[stn_key][ds]['repair_in'] += 1
                            planning[stn_key][ds]['groups_repair_in'][zgroup] = planning[stn_key][ds]['groups_repair_in'].get(zgroup, 0) + 1
                            planning[stn_key][ds]['repair_in_details'].append(dict(res_detail, type='res'))
                        elif is_lt:
                            planning[stn_key][ds]['res_in_lt'] += 1
                            planning[stn_key][ds]['groups_in_lt'][zgroup] = planning[stn_key][ds]['groups_in_lt'].get(zgroup, 0) + 1
                            planning[stn_key][ds]['res_in_details'].append(res_detail)
                        else:
                            planning[stn_key][ds]['res_in'] += 1
                            planning[stn_key][ds]['groups_in'][zgroup] = planning[stn_key][ds]['groups_in'].get(zgroup, 0) + 1
                            planning[stn_key][ds]['res_in_details'].append(res_detail)
        pools_data = {}
        for pool in sorted(all_pools_set):
            _ensure_pool(pool)
            avail_today = avail_real_by_pool.get(pool, 0)
            avail_total = sum(avail_real_by_pool_group.get(pool, {}).values())
            fleet_total = fleet_total_by_pool.get(pool, 0)
            pool_unknown = unknown_status_by_pool.get(pool, 0)
            running_balance = avail_today
            running_fleet_total = fleet_total
            fixed_parados = fleet_oficina_by_pool.get(pool, 0) + fleet_pronto_of_by_pool.get(pool, 0) + blocked_count_by_pool.get(pool, 0)
            fixed_impro = impro_by_pool.get(pool, 0)
            running_grp = dict(avail_real_by_pool_group.get(pool, {}))
            rows = []
            for d in date_range:
                ds = d.isoformat()
                day = planning[pool].get(ds, {'res_out': 0, 'res_out_lt': 0, 'res_in': 0, 'res_in_lt': 0, 'fleet_in': 0, 'fleet_in_lt': 0, 'defleet_out': 0, 'groups_out': {}, 'groups_out_lt': {}, 'groups_in': {}, 'groups_in_lt': {}, 'groups_fleet_in': {}, 'groups_fleet_in_lt': {}, 'groups_defleet_out': {}, 'res_out_details': [], 'res_in_details': [], 'fleet_in_details': [], 'defleet_details': []})
                eff_in = day['fleet_in'] + day['res_in']
                running_balance -= day['defleet_out']
                running_fleet_total -= day['defleet_out']
                for g, c in day.get('groups_defleet_out', {}).items():
                    running_grp[g] = running_grp.get(g, 0) - c
                day_forecast = {g: c for g, c in running_grp.items() if c > 0}
                day_occ_base = max(running_fleet_total - pool_unknown, 0)
                day_occ_pct = round((day_occ_base - fixed_parados - fixed_impro - running_balance) / day_occ_base * 100, 1) if day_occ_base > 0 else None
                day_impro_pct = round(fixed_impro / day_occ_base * 100, 1) if day_occ_base > 0 else None
                rows.append({'date': ds, 'fleet_total': fleet_total, 'available': avail_today if d == today else None, 'saldo': running_balance, 'occupancy_pct': day_occ_pct, 'impro_pct': day_impro_pct, 'forecast_groups': day_forecast, **({'assistant_group_balances': dict(running_grp)} if _assistant else {}), 'res_in': day['res_in'], 'res_out': day['res_out'], 'res_out_lt': day['res_out_lt'], 'res_in_lt': day['res_in_lt'], 'fleet_in': day['fleet_in'], 'fleet_in_lt': day['fleet_in_lt'], 'defleet_out': day['defleet_out'], 'rented_defleet': day.get('rented_defleet', 0), 'groups_out': day['groups_out'], 'groups_out_lt': day['groups_out_lt'], 'groups_in': day['groups_in'], 'groups_in_lt': day['groups_in_lt'], 'groups_fleet_in': day.get('groups_fleet_in', {}), 'groups_fleet_in_lt': day.get('groups_fleet_in_lt', {}), 'groups_defleet_out': day.get('groups_defleet_out', {}), 'groups_rented_defleet': day.get('groups_rented_defleet', {}), 'repair_in': day.get('repair_in', 0), 'groups_repair_in': day.get('groups_repair_in', {}), 'repair_in_details': day.get('repair_in_details', []), 'res_out_details': day.get('res_out_details', []), 'res_in_details': day.get('res_in_details', []), 'fleet_in_details': day.get('fleet_in_details', []), 'defleet_details': day.get('defleet_details', [])})
                running_balance = running_balance + eff_in - day['res_out']
                running_fleet_total += day['fleet_in']
                for g, c in day.get('groups_fleet_in', {}).items():
                    running_grp[g] = running_grp.get(g, 0) + c
                for g, c in day.get('groups_in', {}).items():
                    running_grp[g] = running_grp.get(g, 0) + c
                for g, c in day.get('groups_out', {}).items():
                    running_grp[g] = running_grp.get(g, 0) - c
            pool_occupancy = _occupancy_calc(fleet_total, pool_unknown, avail_today, fleet_oficina_by_pool.get(pool, 0), fleet_pronto_of_by_pool.get(pool, 0), blocked_count_by_pool.get(pool, 0), impro_by_pool.get(pool, 0))
            if rows:
                rows[0]['occupancy_pct'] = pool_occupancy['pct']
                rows[0]['impro_pct'] = pool_occupancy['impro_pct']
            pools_data[pool] = {'rows': rows, 'fleet_total': fleet_total, 'available_today': avail_today, 'available_total': avail_total, 'oficina_count': fleet_oficina_by_pool.get(pool, 0), 'groups_oficina': fleet_oficina_by_pool_group.get(pool, {}), 'pronto_of_count': fleet_pronto_of_by_pool.get(pool, 0), 'groups_pronto_of': fleet_pronto_of_by_pool_group.get(pool, {}), 'groups_available': avail_real_by_pool_group.get(pool, {}), 'avail_vehicles': avail_vehicles_by_pool.get(pool, []), 'oficina_vehicles': oficina_vehicles_by_pool.get(pool, []), 'pronto_of_vehicles': pronto_of_vehicles_by_pool.get(pool, []), 'blocked_count': blocked_count_by_pool.get(pool, 0), 'blocked_groups': blocked_groups_by_pool.get(pool, {}), 'blocked_vehicles': blocked_vehicles_by_pool.get(pool, []), 'no_show_count': len(noshow_by_pool.get(pool, [])), 'groups_no_show': noshow_groups_by_pool.get(pool, {}), 'no_show_details': noshow_by_pool.get(pool, []), 'overdue_ret_count': len(overdue_ret_by_pool.get(pool, [])), 'groups_overdue_ret': overdue_ret_groups_by_pool.get(pool, {}), 'overdue_ret_details': overdue_ret_by_pool.get(pool, []), 'unknown_status_count': pool_unknown, 'impro_count': impro_by_pool.get(pool, 0), 'groups_impro': impro_by_pool_group.get(pool, {}), 'impro_vehicles': impro_vehicles_by_pool.get(pool, []), 'occupancy': pool_occupancy, 'occupancy_by_group': {g: _occupancy_calc(fleet_total_by_pool_group.get(pool, {}).get(g, 0), unknown_status_by_pool_group.get(pool, {}).get(g, 0), avail_real_by_pool_group.get(pool, {}).get(g, 0), fleet_oficina_by_pool_group.get(pool, {}).get(g, 0), fleet_pronto_of_by_pool_group.get(pool, {}).get(g, 0), blocked_groups_by_pool.get(pool, {}).get(g, 0), impro_by_pool_group.get(pool, {}).get(g, 0)) for g in fleet_total_by_pool_group.get(pool, {})}}
        stations_data = {}
        _day_template = {'res_out': 0, 'res_out_lt': 0, 'res_in': 0, 'res_in_lt': 0, 'fleet_in': 0, 'fleet_in_lt': 0, 'defleet_out': 0, 'repair_in': 0, 'groups_repair_in': {}, 'repair_in_details': [], 'groups_out': {}, 'groups_out_lt': {}, 'groups_in': {}, 'groups_in_lt': {}, 'groups_fleet_in': {}, 'groups_fleet_in_lt': {}, 'groups_defleet_out': {}, 'res_out_details': [], 'res_in_details': [], 'fleet_in_details': [], 'defleet_details': []}
        for sid in sorted(all_stations_set):
            stn_key = _STN + sid
            avail_today_s = avail_real_by_stn.get(sid, 0)
            fleet_total_s = fleet_total_by_stn.get(sid, 0)
            running_s = avail_today_s
            running_grp_s = dict(avail_real_by_stn_group.get(sid, {}))
            stn_rows = []
            for d in date_range:
                ds = d.isoformat()
                day = planning.get(stn_key, {}).get(ds, _day_template)
                eff_in_s = day['fleet_in'] + day['res_in']
                running_s -= day['defleet_out']
                for g, c in day.get('groups_defleet_out', {}).items():
                    running_grp_s[g] = running_grp_s.get(g, 0) - c
                day_forecast_s = {g: c for g, c in running_grp_s.items() if c > 0}
                stn_rows.append({'date': ds, 'fleet_total': fleet_total_s, 'available': avail_today_s if d == today else None, 'saldo': running_s, 'forecast_groups': day_forecast_s, 'res_in': day['res_in'], 'res_out': day['res_out'], 'res_out_lt': day['res_out_lt'], 'res_in_lt': day['res_in_lt'], 'fleet_in': day['fleet_in'], 'fleet_in_lt': day['fleet_in_lt'], 'defleet_out': day['defleet_out'], 'rented_defleet': day.get('rented_defleet', 0), 'groups_out': day['groups_out'], 'groups_out_lt': day['groups_out_lt'], 'groups_in': day['groups_in'], 'groups_in_lt': day['groups_in_lt'], 'groups_fleet_in': day.get('groups_fleet_in', {}), 'groups_fleet_in_lt': day.get('groups_fleet_in_lt', {}), 'groups_defleet_out': day.get('groups_defleet_out', {}), 'groups_rented_defleet': day.get('groups_rented_defleet', {}), 'repair_in': day.get('repair_in', 0), 'groups_repair_in': day.get('groups_repair_in', {}), 'repair_in_details': day.get('repair_in_details', []), 'res_out_details': day.get('res_out_details', []), 'res_in_details': day.get('res_in_details', []), 'fleet_in_details': day.get('fleet_in_details', []), 'defleet_details': day.get('defleet_details', [])})
                running_s = running_s + eff_in_s - day['res_out']
                for g, c in day.get('groups_fleet_in', {}).items():
                    running_grp_s[g] = running_grp_s.get(g, 0) + c
                for g, c in day.get('groups_in', {}).items():
                    running_grp_s[g] = running_grp_s.get(g, 0) + c
                for g, c in day.get('groups_out', {}).items():
                    running_grp_s[g] = running_grp_s.get(g, 0) - c
            stations_data[sid] = {'rows': stn_rows, 'fleet_total': fleet_total_s, 'available_today': avail_today_s, 'available_total': sum(avail_real_by_stn_group.get(sid, {}).values()), 'oficina_count': fleet_oficina_by_stn.get(sid, 0), 'groups_oficina': fleet_oficina_by_stn_group.get(sid, {}), 'pronto_of_count': fleet_pronto_of_by_stn.get(sid, 0), 'groups_pronto_of': fleet_pronto_of_by_stn_group.get(sid, {}), 'groups_available': avail_real_by_stn_group.get(sid, {}), 'avail_vehicles': avail_vehicles_by_stn.get(sid, []), 'oficina_vehicles': oficina_vehicles_by_stn.get(sid, []), 'pronto_of_vehicles': pronto_of_vehicles_by_stn.get(sid, []), 'blocked_count': blocked_count_by_stn.get(sid, 0), 'blocked_groups': blocked_groups_by_stn.get(sid, {}), 'blocked_vehicles': blocked_vehicles_by_stn.get(sid, []), 'no_show_count': len(noshow_by_stn.get(sid, [])), 'no_show_details': noshow_by_stn.get(sid, []), 'impro_count': len(impro_vehicles_by_stn.get(sid, [])), 'impro_vehicles': impro_vehicles_by_stn.get(sid, []), **({'occupancy_by_group': occupancy_counts_by_stn.get(sid, {}), 'unknown_status_count': sum((x['unknown'] for x in occupancy_counts_by_stn.get(sid, {}).values())), 'name': station_names.get(sid, sid), 'station_ids': [sid]} if _assistant else {})}
        assistant_stations = dict(stations_data) if _assistant else None
        _merge_done = set()
        _actually_merged_secondary = set()
        for sid_a, sid_b in station_paired.items():
            pair_key = tuple(sorted([sid_a, sid_b]))
            if pair_key in _merge_done:
                continue
            _merge_done.add(pair_key)
            if station_tipos.get(pair_key[0]) != 'Station' or station_tipos.get(pair_key[1]) != 'Station':
                continue
            sid_primary, sid_secondary = pair_key
            if sid_primary not in stations_data and sid_secondary not in stations_data:
                continue
            data_a = stations_data.get(sid_primary)
            data_b = stations_data.get(sid_secondary)
            if not data_a and (not data_b):
                continue
            name_a = station_names.get(sid_primary, sid_primary)
            name_b = station_names.get(sid_secondary, sid_secondary)
            merged_name = name_a + ' + ' + name_b

            def _merge_dicts(d1, d2):
                r = dict(d1)
                for k, v in d2.items():
                    r[k] = r.get(k, 0) + v
                return r
            if data_a and data_b:
                m = {'fleet_total': data_a['fleet_total'] + data_b['fleet_total'], 'available_today': data_a['available_today'] + data_b['available_today'], 'available_total': data_a['available_total'] + data_b['available_total'], 'oficina_count': data_a['oficina_count'] + data_b['oficina_count'], 'pronto_of_count': data_a.get('pronto_of_count', 0) + data_b.get('pronto_of_count', 0), 'blocked_count': data_a['blocked_count'] + data_b['blocked_count'], 'groups_oficina': _merge_dicts(data_a['groups_oficina'], data_b['groups_oficina']), 'groups_pronto_of': _merge_dicts(data_a.get('groups_pronto_of', {}), data_b.get('groups_pronto_of', {})), 'groups_available': _merge_dicts(data_a['groups_available'], data_b['groups_available']), 'groups_blocked': _merge_dicts(data_a.get('blocked_groups', {}), data_b.get('blocked_groups', {})), 'blocked_groups': _merge_dicts(data_a.get('blocked_groups', {}), data_b.get('blocked_groups', {})), 'avail_vehicles': data_a['avail_vehicles'] + data_b['avail_vehicles'], 'oficina_vehicles': data_a['oficina_vehicles'] + data_b['oficina_vehicles'], 'pronto_of_vehicles': data_a.get('pronto_of_vehicles', []) + data_b.get('pronto_of_vehicles', []), 'blocked_vehicles': data_a['blocked_vehicles'] + data_b['blocked_vehicles'], 'no_show_count': data_a.get('no_show_count', 0) + data_b.get('no_show_count', 0), 'no_show_details': data_a.get('no_show_details', []) + data_b.get('no_show_details', [])}
                rows_a = data_a['rows']
                rows_b = data_b['rows']
                merged_rows = []
                for ri in range(max(len(rows_a), len(rows_b))):
                    ra = rows_a[ri] if ri < len(rows_a) else {}
                    rb = rows_b[ri] if ri < len(rows_b) else {}
                    mr = {'date': ra.get('date') or rb.get('date'), 'fleet_total': ra.get('fleet_total', 0) + rb.get('fleet_total', 0), 'available': (ra.get('available') or 0) + (rb.get('available') or 0) if ra.get('available') is not None or rb.get('available') is not None else None, 'saldo': ra.get('saldo', 0) + rb.get('saldo', 0), 'forecast_groups': _merge_dicts(ra.get('forecast_groups', {}), rb.get('forecast_groups', {})), 'res_in': ra.get('res_in', 0) + rb.get('res_in', 0), 'res_out': ra.get('res_out', 0) + rb.get('res_out', 0), 'res_out_lt': ra.get('res_out_lt', 0) + rb.get('res_out_lt', 0), 'res_in_lt': ra.get('res_in_lt', 0) + rb.get('res_in_lt', 0), 'fleet_in': ra.get('fleet_in', 0) + rb.get('fleet_in', 0), 'fleet_in_lt': ra.get('fleet_in_lt', 0) + rb.get('fleet_in_lt', 0), 'defleet_out': ra.get('defleet_out', 0) + rb.get('defleet_out', 0), 'groups_out': _merge_dicts(ra.get('groups_out', {}), rb.get('groups_out', {})), 'groups_out_lt': _merge_dicts(ra.get('groups_out_lt', {}), rb.get('groups_out_lt', {})), 'groups_in': _merge_dicts(ra.get('groups_in', {}), rb.get('groups_in', {})), 'groups_in_lt': _merge_dicts(ra.get('groups_in_lt', {}), rb.get('groups_in_lt', {})), 'groups_fleet_in': _merge_dicts(ra.get('groups_fleet_in', {}), rb.get('groups_fleet_in', {})), 'groups_fleet_in_lt': _merge_dicts(ra.get('groups_fleet_in_lt', {}), rb.get('groups_fleet_in_lt', {})), 'groups_defleet_out': _merge_dicts(ra.get('groups_defleet_out', {}), rb.get('groups_defleet_out', {})), 'res_out_details': ra.get('res_out_details', []) + rb.get('res_out_details', []), 'res_in_details': ra.get('res_in_details', []) + rb.get('res_in_details', []), 'fleet_in_details': ra.get('fleet_in_details', []) + rb.get('fleet_in_details', []), 'defleet_details': ra.get('defleet_details', []) + rb.get('defleet_details', []), 'repair_in': ra.get('repair_in', 0) + rb.get('repair_in', 0), 'groups_repair_in': _merge_dicts(ra.get('groups_repair_in', {}), rb.get('groups_repair_in', {})), 'repair_in_details': ra.get('repair_in_details', []) + rb.get('repair_in_details', [])}
                    merged_rows.append(mr)
                m['rows'] = merged_rows
            else:
                m = dict(data_a or data_b)
            stations_data[sid_primary] = m
            if sid_secondary in stations_data:
                del stations_data[sid_secondary]
            station_names[sid_primary] = merged_name
            _actually_merged_secondary.add(sid_secondary)
        pool_stations_map = {}
        for sid, pname in station_pool.items():
            if pool_filter and pname != pool_filter:
                continue
            if sid in _actually_merged_secondary:
                continue
            pool_stations_map.setdefault(pname, [])
            pool_stations_map[pname].append({'id': sid, 'nome': station_names.get(sid, sid), 'tipo': station_tipos.get(sid, '')})
        for p in pool_stations_map:
            pool_stations_map[p].sort(key=lambda s: s['nome'])
        all_zgroups_raw = set()
        for r in all_res:
            zg = r.get('ZGROUP', '')
            if zg:
                all_zgroups_raw.add(acriss_aliases.get(zg, zg))
        for pool_grp in fleet_by_pool_group.values():
            all_zgroups_raw.update(pool_grp.keys())
        for pk in planning:
            for ds_data in planning[pk].values():
                all_zgroups_raw.update(ds_data.get('groups_fleet_in', {}).keys())
                all_zgroups_raw.update(ds_data.get('groups_fleet_in_lt', {}).keys())
        all_zgroups = sorted(all_zgroups_raw)
        planning_cfg = {}
        if os.path.exists(PLANNING_CONFIG_FILE):
            try:
                with open(PLANNING_CONFIG_FILE, 'r', encoding='utf-8') as f:
                    planning_cfg = json.load(f)
            except Exception:
                pass
        return jsonify({'pools': pools_data, 'stations': stations_data, **({'assistant_stations': assistant_stations, 'assistant_pool_members': {pool: [sid for sid, name in station_pool.items() if name == pool] for pool in pools_data}} if _assistant else {}), 'pool_stations': pool_stations_map, 'planning_config': planning_cfg, 'pool_names': sorted(all_pools_set), 'zgroups': all_zgroups, 'acriss_categories': [{'nome': c.get('nome', ''), 'groups': c.get('groups', [])} for c in acriss_categories], 'planning_regions': [{'nome': r.get('nome', ''), 'pools': r.get('pools', [])} for r in planning_regions], 'pool_folders': [{'nome': f.get('nome', ''), 'icon': f.get('icon', ''), 'pools': f.get('pools', [])} for f in pool_folders], 'days': days, 'date_from': today.isoformat(), 'date_to': end_date.isoformat(), 'fleet_captured_at': fleet_captured, 'res_fetched_at': res_fetched_at, 'res_total_cache': res_total_cache, 'res_counted_planning': res_counted, 'infleet': _infleet_cache.get('data'), 'infleet_fetch_status': {} if _assistant else _infleet_fetch_status()})
    except Exception as e:
        logger.error(f'Fleet planning error: {e}', exc_info=True)
        return (jsonify({'error': str(e)}), 500)

def _build_inverse_upgrade_map(upgrade_matrix):
    inverse = {}
    all_groups = set(upgrade_matrix.keys())
    for grp in all_groups:
        inverse.setdefault(grp, [(grp, 0)])
    for res_grp, tiers in upgrade_matrix.items():
        for veh_grp in tiers.get('preferencial', []):
            inverse.setdefault(veh_grp, [(veh_grp, 0)])
            inverse[veh_grp].append((res_grp, 1))
        for veh_grp in tiers.get('aceitavel', []):
            inverse.setdefault(veh_grp, [(veh_grp, 0)])
            inverse[veh_grp].append((res_grp, 2))
        for veh_grp in tiers.get('evitar', []):
            inverse.setdefault(veh_grp, [(veh_grp, 0)])
            inverse[veh_grp].append((res_grp, 3))
    for grp in inverse:
        inverse[grp] = sorted(set(inverse[grp]), key=lambda x: x[1])
    return inverse

def _match_vehicle_by_plate(plate: str, *, read_only=False) -> dict:
    plate = plate.strip().upper()
    tipo_map = _read_fleet_type_map(read_only=read_only)
    fleet_data = _not_avail_fleet_cache.get('data') or []
    zc = _get_zone_maps()
    id_zona = zc['id_zona']
    id_region = zc['id_region']
    planning_cfg = {}
    try:
        with open(PLANNING_CONFIG_FILE, 'r', encoding='utf-8') as f:
            planning_cfg = json.load(f)
    except Exception:
        pass
    upgrade_matrix = planning_cfg.get('upgrade_matrix', {})
    inverse_map = _build_inverse_upgrade_map(upgrade_matrix)
    acriss_aliases: dict = {}
    station_names: dict = {}
    try:
        scfg = _stations_config()
        acriss_aliases = scfg.get('acriss_aliases', {})
        for sid, sdata in scfg.get('stations', {}).items():
            station_names[sid] = sdata.get('nome', sid)
    except Exception:
        pass
    vehicle = None
    for v in fleet_data:
        if v.get('licensePlate', '').strip().upper() == plate:
            vehicle = v
            break
    if vehicle is None:
        return {'error': f'Vehicle {plate} was not found in the fleet snapshot. Check the fleet data.'}
    branch = _resolve_vehicle_branch(vehicle, set(station_names.keys()))
    acriss = str(vehicle.get('acrissCode', '') or '')[:4].strip().upper()
    acriss = acriss_aliases.get(acriss, acriss)
    cs = vehicle.get('classicStatus', -1)
    is_rented = cs == 0
    is_available = cs in (1, 2, 3, 4, 7, 9, 10, 11)
    now = datetime.now()
    avail_datetime = now
    eff_station_id = branch
    eff_pool = id_zona.get(branch, '')
    eff_region = id_region.get(branch, '')
    avail_label = 'Already available'
    if is_rented:
        ret_time_str = vehicle.get('returnTime', '') or ''
        if ret_time_str:
            try:
                avail_datetime = datetime.strptime(ret_time_str[:16], '%Y-%m-%dT%H:%M')
                avail_label = f"Return {avail_datetime.strftime('%H:%M / %d/%m/%Y')}"
                ret_branch = str((vehicle.get('returnBranch') or {}).get('id', '') if isinstance(vehicle.get('returnBranch'), dict) else vehicle.get('returnBranch', '') or '')
                if ret_branch:
                    eff_station_id = ret_branch
                    eff_pool = id_zona.get(ret_branch, eff_pool)
                    eff_region = id_region.get(ret_branch, eff_region)
            except (ValueError, TypeError):
                pass
    tipo_info = tipo_map.get(plate)
    tipo = tipo_info['tipo_compra'] if tipo_info else None
    tipo_label = {'FP': 'Owned Fleet', 'BB': 'Buy Back'}.get(tipo, 'Unknown') if tipo else 'No ownership data'
    contract_end_dt = None
    contract_end_str = ''
    remaining_days = None
    raw_end = vehicle.get('defleetDate', '') or ''
    if raw_end:
        try:
            contract_end_dt = datetime.strptime(raw_end[:10], '%Y-%m-%d').replace(hour=23, minute=59)
            contract_end_str = contract_end_dt.strftime('%d/%m/%Y')
            remaining_days = (contract_end_dt.date() - date.today()).days
        except (ValueError, TypeError):
            pass
    vehicle_info = {'plate': plate, 'group': acriss, 'name': vehicle.get('displayName', ''), 'engine': vehicle.get('engineType', '') or '', 'tipo': tipo or '?', 'tipo_label': tipo_label, 'station_id': branch, 'station': station_names.get(branch, branch), 'pool': id_zona.get(branch, ''), 'region': id_region.get(branch, ''), 'eff_station_id': eff_station_id, 'eff_station': station_names.get(eff_station_id, eff_station_id), 'eff_pool': eff_pool, 'eff_region': eff_region, 'status': 'rented' if is_rented else 'available' if is_available else 'other', 'status_code': cs, 'remaining_days': remaining_days, 'contract_end': contract_end_str, 'avail_label': avail_label}
    all_res = _res_bulk_cache.get('rows') or []
    if not all_res:
        return {'vehicle': vehicle_info, 'suggestions': [], 'no_match_reason': 'No cached reservations available.', 'nearby_alternatives': []}
    if contract_end_dt is None:
        contract_end_dt = now + timedelta(days=30)
    MAX_EXTENSION_DAYS = 5
    if contract_end_dt >= now:
        horizon_dt = min(contract_end_dt + timedelta(days=MAX_EXTENSION_DAYS), now + timedelta(days=30))
    else:
        horizon_dt = now + timedelta(days=30)
    target_region = 'Central' if tipo == 'BB' else eff_region
    compat_list = inverse_map.get(acriss, [(acriss, 0)])
    compat_groups = {g for g, _ in compat_list}
    SAME_STATION_BUFFER_H = 2
    VALID_STATUSES = {'RS', 'RQ'}
    all_candidates = []
    already_blocked = []
    for r in all_res:
        if r.get('STATUS') not in VALID_STATUSES:
            continue
        if r.get('IS_LONGTERM'):
            continue
        pick_dt_str = r.get('PICK_DATETIME', '') or ''
        if not pick_dt_str:
            continue
        try:
            pick_dt = datetime.fromisoformat(pick_dt_str[:19])
        except (ValueError, TypeError):
            continue
        if pick_dt < now or pick_dt > horizon_dt:
            continue
        ret_dt_str = r.get('RET_DATETIME', '') or ''
        try:
            ret_dt_parsed = datetime.fromisoformat(ret_dt_str[:19]) if ret_dt_str else None
        except (ValueError, TypeError):
            ret_dt_parsed = None
        if ret_dt_parsed and contract_end_dt >= now:
            if ret_dt_parsed > contract_end_dt + timedelta(days=MAX_EXTENSION_DAYS):
                continue
        zgroup = str(r.get('ZGROUP', '') or '').strip().upper()
        zgroup = acriss_aliases.get(zgroup, zgroup)
        if zgroup not in compat_groups:
            continue
        pick_sid = str(r.get('PICK_STATION_ID', '') or '')
        ret_sid = str(r.get('RET_STATION_ID', '') or '')
        pick_pool = id_zona.get(pick_sid, '')
        pick_region = id_region.get(pick_sid, '')
        ret_region = id_region.get(ret_sid, '')
        one_way = bool(r.get('ONE_WAY'))
        if tipo == 'BB':
            if not one_way or ret_region != 'Central':
                continue
        if pick_sid == eff_station_id:
            if pick_dt < avail_datetime + timedelta(hours=SAME_STATION_BUFFER_H):
                continue
            loc_priority = 0
        elif pick_pool == eff_pool and eff_pool:
            if pick_dt < avail_datetime + timedelta(hours=24):
                continue
            loc_priority = 1
        elif pick_region == eff_region and eff_region:
            if pick_dt < avail_datetime + timedelta(hours=24):
                continue
            loc_priority = 2
        else:
            continue
        group_priority = next((p for g, p in compat_list if g == zgroup), 3)
        assigned_plate = str(r.get('LICENCE_PLATE', '') or '').strip().upper()
        if assigned_plate:
            already_blocked.append({'res_no': r.get('RESERVATION_NO', ''), 'assigned_plate': assigned_plate, 'group': zgroup, 'pick_station': r.get('PICK_STATION_NAME') or station_names.get(pick_sid, pick_sid), 'pick_time': pick_dt.strftime('%H:%M / %d/%m/%Y'), 'ret_station': r.get('RET_STATION_NAME') or station_names.get(ret_sid, ret_sid), 'ret_time': ret_dt_parsed.strftime('%H:%M / %d/%m/%Y') if ret_dt_parsed else '', 'loc_match': ['Same Station', 'Mesma Pool', 'Mesma Region'][loc_priority], '_sort_ts': pick_dt.timestamp()})
            continue
        if one_way:
            direction = 'outbound' if ret_region != eff_region else 'local_ow'
        else:
            direction = 'roundtrip'
        extension_days = 0
        if ret_dt_parsed:
            if remaining_days is not None:
                effective_end = now + timedelta(days=max(0, remaining_days - 1))
            else:
                effective_end = contract_end_dt
            if ret_dt_parsed > effective_end:
                extension_days = int((ret_dt_parsed - effective_end).days)
        if extension_days > MAX_EXTENSION_DAYS:
            continue
        if extension_days > 0 and remaining_days is not None and (remaining_days <= 7):
            continue
        all_candidates.append({'res_no': r.get('RESERVATION_NO', ''), 'group': zgroup, 'match_type': ['Exato', 'Preferencial', 'Acceptable', 'Evitar'][min(group_priority, 3)], 'loc_match': ['Same Station', 'Mesma Pool', 'Mesma Region'][loc_priority], 'pick_station': r.get('PICK_STATION_NAME') or station_names.get(pick_sid, pick_sid), 'pick_station_id': pick_sid, 'pick_time': pick_dt.strftime('%H:%M / %d/%m/%Y'), 'ret_station': r.get('RET_STATION_NAME') or station_names.get(ret_sid, ret_sid), 'ret_station_id': ret_sid, 'ret_region': ret_region, 'ret_time': ret_dt_parsed.strftime('%H:%M / %d/%m/%Y') if ret_dt_parsed else '', 'duration': r.get('DURATION', 0) or 0, 'one_way': one_way, 'direction': direction, 'needs_extension': extension_days > 0, 'extension_days': extension_days, '_sort': (loc_priority, extension_days > 0, group_priority, pick_dt.timestamp())})
    all_candidates.sort(key=lambda x: x['_sort'])
    for c in all_candidates:
        del c['_sort']
    already_blocked.sort(key=lambda x: x.get('_sort_ts', 0))
    for c in already_blocked:
        c.pop('_sort_ts', None)
    suggestions = all_candidates[:5]
    no_match_reason = None
    nearby_alternatives = []
    if not suggestions:
        horizon_d = min(int((horizon_dt - now).days), 30)
        if tipo == 'BB':
            no_match_reason = f"No one-way reservations from {vehicle_info['eff_station']} to Lisbon/Central within the next {horizon_d} days."
        else:
            no_match_reason = f"No available reservations at {vehicle_info['eff_station']} ({vehicle_info['eff_pool']}) within the next {horizon_d} days."
        seen_alt: dict[str, int] = {}
        for r in all_res:
            if r.get('STATUS') not in VALID_STATUSES:
                continue
            if r.get('LICENCE_PLATE') or r.get('IS_LONGTERM'):
                continue
            pick_sid_a = str(r.get('PICK_STATION_ID', '') or '')
            if not pick_sid_a or pick_sid_a == eff_station_id:
                continue
            pick_pool_a = id_zona.get(pick_sid_a, '')
            pick_region_a = id_region.get(pick_sid_a, '')
            zgroup_a = acriss_aliases.get(str(r.get('ZGROUP', '') or '').strip().upper(), str(r.get('ZGROUP', '') or '').strip().upper())
            if zgroup_a not in compat_groups:
                continue
            if tipo == 'BB':
                ret_sid_a = str(r.get('RET_STATION_ID', '') or '')
                if id_region.get(ret_sid_a, '') != 'Central' or not r.get('ONE_WAY'):
                    continue
                if pick_pool_a != eff_pool and pick_region_a != eff_region:
                    continue
            elif pick_pool_a != eff_pool and pick_region_a != eff_region:
                continue
            seen_alt[pick_sid_a] = seen_alt.get(pick_sid_a, 0) + 1
        for sid_a, cnt in sorted(seen_alt.items(), key=lambda x: -x[1])[:5]:
            nearby_alternatives.append({'station_id': sid_a, 'station': station_names.get(sid_a, sid_a), 'pool': id_zona.get(sid_a, ''), 'res_count': cnt})
    return {'vehicle': vehicle_info, 'suggestions': suggestions, 'already_blocked': already_blocked[:8], 'no_match_reason': no_match_reason, 'nearby_alternatives': nearby_alternatives}

@app.route('/api/island-blocking')
def api_island_blocking(_assistant=False):
    import re as _re
    from demo_clock import DemoDateTime as _dt
    zone_filter = request.args.get('zone', '').strip()
    groups_raw = request.args.get('groups', '').strip()
    tipo_filter = request.args.get('tipo', '').strip().upper()
    min_days = int(request.args.get('min_days', 200))
    groups_set = set((g.strip().upper() for g in groups_raw.split(',') if g.strip())) if groups_raw else None
    try:
        naf_vehs = _not_avail_fleet_cache.get('data') or []
        if not naf_vehs:
            with open(CACHE_FILE, 'r', encoding='utf-8') as f:
                cache = json.load(f)
            naf_vehs = cache.get('vehicles', [])
        maps = _get_zone_maps()
        id_zona = maps['id_zona']
        id_region = maps['id_region']
        station_names = {}
        try:
            _st_cfg = _stations_config()
            for sid, sdata in _st_cfg.get('stations', {}).items():
                station_names[sid] = sdata.get('nome', sid)
        except Exception:
            pass
        tipo_map = {}
        try:
            tipo_map = _read_fleet_type_map(read_only=_assistant)
        except Exception:
            pass
        if not tipo_map:
            try:
                with open(os.path.join(DATA_DIR, 'infleet_cache.json'), 'r', encoding='utf-8') as f:
                    inf_cache = json.load(f)
                for item in inf_cache:
                    p = item.get('plate', '').strip().upper()
                    if p:
                        tipo_map[p] = {'tipo_compra': item.get('tipo_compra', '')}
            except Exception:
                pass
        today = _dt.now().replace(hour=0, minute=0, second=0, microsecond=0)
        pt_plate = _re.compile('^[A-Z]{2}-?\\d{2}-?[A-Z]{2}$')
        _SERVICE_RATES = {'PTL00000', 'PTL34000', 'PTL35000', 'PTL42000', 'PTL40000', 'PTL36000', 'PTL30000', 'PTL41000'}
        avail_statuses = {1, 2, 3, 4, 7, 9, 10, 11}
        valid_statuses = {0, 1, 2, 3, 4, 7, 9, 10, 11}
        results = []
        for v in naf_vehs:
            cs = v.get('classicStatus', -1)
            if cs not in valid_statuses:
                continue
            is_available = cs in avail_statuses
            is_rented = cs == 0
            plate = (v.get('licensePlate', '') or '').strip().upper()
            if not pt_plate.match(plate):
                continue
            ra = v.get('rentalActivity') or {}
            if ra.get('rateCode', '') in _SERVICE_RATES:
                continue
            if ra.get('isLongterm'):
                continue
            _hold_desc = v.get('holdDescriptionForDueDate') or ''
            _hold_code = v.get('holdCodeForDueDate') or ''
            if _hold_desc or _hold_code:
                continue
            defleet_str = v.get('defleetDate', '') or ''
            if not defleet_str:
                continue
            try:
                end_date = _dt.strptime(defleet_str[:10], '%Y-%m-%d')
            except (ValueError, TypeError):
                continue
            remaining_days = (end_date - today).days
            if remaining_days < min_days:
                continue
            branch = _resolve_vehicle_branch(v, set(id_region.keys()))
            pool = id_zona.get(branch, '')
            region = id_region.get(branch, '')
            station = station_names.get(branch, branch)
            if zone_filter:
                if zone_filter != pool and zone_filter != region:
                    continue
            acriss = (v.get('acrissCode') or '').strip().upper()
            if groups_set and acriss not in groups_set:
                continue
            tipo = (tipo_map.get(plate, {}).get('tipo_compra', '') or '').upper()
            if not tipo:
                tipo = '?'
            if tipo_filter and tipo != tipo_filter:
                continue
            hold_remaining = v.get('holdRemainingMileage')
            if hold_remaining is None:
                hold_remaining = v.get('remainingMileage')
            if hold_remaining is not None and hold_remaining <= 10:
                continue
            ret_time = ''
            ret_time_fmt = ''
            ret_station_name = ''
            if is_rented:
                ret_time = v.get('returnTime', '') or ''
                if ret_time:
                    try:
                        _ret_dt = _dt.strptime(ret_time[:16], '%Y-%m-%dT%H:%M')
                        if _ret_dt < today:
                            continue
                        ret_time_fmt = _ret_dt.strftime('%H:%M / %d/%m/%Y')
                    except (ValueError, TypeError):
                        pass
                ret_branch = str(v.get('returnBranch', {}).get('id', '') if isinstance(v.get('returnBranch'), dict) else v.get('returnBranch', '') or '')
                _ret_sid = ret_branch if ret_branch else branch
                ret_station_name = station_names.get(_ret_sid, _ret_sid)
            results.append({'plate': plate, 'group': acriss, 'tipo': tipo, 'name': v.get('displayName', ''), 'station': station, 'station_id': branch, 'pool': pool, 'region': region, 'remaining_days': remaining_days, 'contract_end': end_date.strftime('%d/%m/%Y'), 'engine': v.get('engineType', ''), 'mileage': v.get('mileage'), 'remaining_km': hold_remaining, 'sub_status': v.get('subStatus', ''), 'parking': v.get('parkingSlot') or v.get('parkingSlotV2') or v.get('parkingParameterV2') or v.get('parkingParameter') or '', 'is_available': is_available, 'return_time': ret_time_fmt, 'return_station': ret_station_name})

        def _ib_sort_key(x):
            if x['is_available']:
                return (0, '', x['plate'])
            else:
                return (1, x.get('return_time', '') or 'zz', x['plate'])
        results.sort(key=_ib_sort_key)
        all_pools = sorted(set(id_zona.values()))
        all_regions = sorted(set(id_region.values()))
        return jsonify({'success': True, 'vehicles': results, 'total': len(results), 'filters': {'zone': zone_filter, 'groups': list(groups_set) if groups_set else [], 'tipo': tipo_filter, 'min_days': min_days}, 'available_zones': all_pools, 'available_regions': all_regions})
    except Exception as e:
        logger.exception('island-blocking error')
        return (jsonify({'success': False, 'error': str(e)}), 500)

def _rotation_filter(cargas, *, estado=None, utilizador=None, motivo=None, data_inicio=None, data_fim=None, pesquisa=None):
    out = []
    q = (pesquisa or '').strip().lower()
    for c in cargas:
        if estado and c.get('estado') != estado:
            continue
        if utilizador and c.get('utilizador') != utilizador:
            continue
        if motivo and c.get('motivo') != motivo:
            continue
        if data_inicio:
            d = (c.get('data_levantamento') or c.get('data_criacao') or '')[:10]
            if d and d < data_inicio:
                continue
        if data_fim:
            d = (c.get('data_levantamento') or c.get('data_criacao') or '')[:10]
            if d and d > data_fim:
                continue
        if q:
            blob = ' '.join([str(c.get('referencia', '')), str(c.get('zona_origem', '')), str(c.get('zona_destino', '')), str(c.get('estacao_origem_nome', '')), str(c.get('estacao_destino_nome', '')), str(c.get('motivo', '')), str(c.get('utilizador', '')), ' '.join([v.get('matricula') or '' for v in c.get('viaturas') or []]), ' '.join([a.get('transportador') or '' for a in c.get('adjudicacoes') or []])]).lower()
            if q not in blob:
                continue
        out.append(c)
    out.sort(key=lambda x: x.get('data_criacao') or '', reverse=True)
    return out

@app.route('/api/rotation/stats', methods=['GET'])
def api_rotation_stats():
    try:
        cargas = _rotation_load_cargas()
        por_estado = {}
        total_viaturas = 0
        for c in cargas:
            est = c.get('estado') or 'Pending'
            por_estado[est] = por_estado.get(est, 0) + 1
            total_viaturas += c.get('num_viaturas') or len(c.get('viaturas') or [])
        return jsonify({'success': True, 'total_cargas': len(cargas), 'total_viaturas': total_viaturas, 'por_estado': por_estado})
    except Exception as e:
        logger.exception('rotation stats error')
        return (jsonify({'success': False, 'error': str(e)}), 500)

@app.route('/api/rotation/historico', methods=['GET'])
def api_rotation_historico():
    try:
        entries = _rotation_load_historico()
        ref = request.args.get('referencia')
        tipo = request.args.get('tipo')
        try:
            limit = int(request.args.get('limit', '500'))
        except ValueError:
            limit = 500
        if ref:
            entries = [e for e in entries if e.get('referencia') == ref]
        if tipo:
            entries = [e for e in entries if e.get('tipo') == tipo]
        entries.sort(key=lambda e: e.get('timestamp') or '', reverse=True)
        if limit > 0:
            entries = entries[:limit]
        return jsonify({'success': True, 'historico': entries, 'total': len(entries)})
    except Exception as e:
        logger.exception('rotation historico error')
        return (jsonify({'success': False, 'error': str(e)}), 500)

@app.route('/api/capacity-control', methods=['GET'])
def api_capacity_control():
    config = _load_capacity_config()
    stations_cfg = config.get('stations', {})
    days_param = min(int(request.args.get('days', 21) or 21), 90)
    today = datetime.now().date()
    date_range = [(today + timedelta(days=i)).isoformat() for i in range(days_param)]
    date_set = set(date_range)
    res_rows = _res_bulk_cache.get('rows') or []
    _REAL_ST = {'RS', 'CO'}
    slots_meta = {}
    _parsed_slots = {}
    for st_id, st_cfg in stations_cfg.items():
        raw_slots = st_cfg.get('slots', [])
        parsed = []
        meta = []
        oot_until = st_cfg.get('oot_until', '00:00') or '00:00'
        oot_h, oot_m = map(int, oot_until.split(':'))
        for sl in raw_slots:
            fh, fm = map(int, sl['from'].split(':'))
            th, tm = map(int, sl['to'].split(':'))
            key = '{}-{}'.format(sl['from'], sl['to'])
            lbl = sl.get('label', key)
            parsed.append({'key': key, 'from_min': fh * 60 + fm, 'to_min': th * 60 + tm, 'max': sl.get('max')})
            meta.append({'key': key, 'label': lbl, 'max': sl.get('max')})
        if oot_h * 60 + oot_m > 0:
            meta.append({'key': '__oot__', 'label': 'Fora Horas', 'max': None})
        slots_meta[st_id] = meta
        _parsed_slots[st_id] = parsed
    data = {}
    for st_id, st_cfg in stations_cfg.items():
        st_data = {}
        for d in date_range:
            st_data[d] = {'total': 0, 'oot': 0, 'slots': {sl['key']: 0 for sl in slots_meta[st_id]}}
        data[st_id] = st_data
    id_to_key = {}
    for cfg_key, st_cfg in stations_cfg.items():
        for sid in st_cfg.get('station_ids', [cfg_key]):
            id_to_key[str(sid)] = cfg_key
    for r in res_rows:
        if r.get('STATUS') not in _REAL_ST:
            continue
        pick_st_id = str(r.get('PICK_STATION_ID', ''))
        cfg_key = id_to_key.get(pick_st_id)
        if cfg_key is None:
            continue
        pick_dt_str = r.get('PICK_DATETIME', '')
        if not pick_dt_str:
            continue
        try:
            pick_dt = datetime.fromisoformat(pick_dt_str)
            if _TZ_LISBON and pick_dt.tzinfo is not None:
                pick_dt_local = pick_dt.astimezone(_TZ_LISBON).replace(tzinfo=None)
            else:
                pick_dt_local = pick_dt.replace(tzinfo=None)
        except Exception:
            continue
        st_cfg = stations_cfg[cfg_key]
        oot_until = st_cfg.get('oot_until', '00:00') or '00:00'
        oot_h, oot_m = map(int, oot_until.split(':'))
        oot_threshold = oot_h * 60 + oot_m
        pick_min = pick_dt_local.hour * 60 + pick_dt_local.minute
        is_oot = oot_threshold > 0 and pick_min < oot_threshold
        eff_date = (pick_dt_local.date() - timedelta(days=1) if is_oot else pick_dt_local.date()).isoformat()
        if eff_date not in date_set:
            continue
        day_data = data[cfg_key][eff_date]
        day_data['total'] += 1
        if is_oot:
            day_data['oot'] += 1
            if '__oot__' in day_data['slots']:
                day_data['slots']['__oot__'] += 1
        else:
            assigned = False
            for sl in _parsed_slots[cfg_key]:
                to_min_adj = sl['to_min'] if sl['to_min'] > 0 else 24 * 60
                if sl['from_min'] <= pick_min <= to_min_adj:
                    if sl['key'] in day_data['slots']:
                        day_data['slots'][sl['key']] += 1
                    assigned = True
                    break
            if not assigned:
                day_data['slots']['__other__'] = day_data['slots'].get('__other__', 0) + 1
    return jsonify({'stations_config': stations_cfg, 'slots_meta': slots_meta, 'days': date_range, 'data': data})
